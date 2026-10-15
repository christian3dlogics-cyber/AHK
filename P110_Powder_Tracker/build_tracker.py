"""P110 Powder Tracker v20 — workbook generator.

    python build_tracker.py [out.xlsx]          -> empty template (P110_Powder_Tracker_v20_template.xlsx)

build_workbook(data, nrows) is also used by migrate_v19.py (pre-populated v20) and by the tests
(fuzz workbooks use a smaller nrows; formulas are identical, only the last row moves).

Design (spec §2.3): every value is computed once per row in helper columns S onward:
  normalized inputs -> slot state -> removal order -> funnel queue (4 lots) -> funnel consumption
  -> slot draws (longest-loaded first, ties by key) -> draw columns -> 72-column bucket ledger
  -> per-category message columns -> R row check.
Compatible with Excel 2016+ / LibreOffice 7+: no LET/LAMBDA/XLOOKUP/FILTER/TEXTJOIN/IFS/MAXIFS,
no dynamic arrays, no volatile functions, no macros.
"""
import functools
import re
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

FIRST = 5                       # first data row on the machine sheets
MACHINES = ("Old P110", "New P110")
KEYS = ["M%d%s-" % (c, b) for c in range(1, 10) for b in "ABCDEFGH"]   # 72 keys, Bucket Status order
BS_FIRST, BS_LAST = 4, 75
SETTINGS_DEFAULTS = {"B3": 1, "B4": 0.6, "B5": 0.6, "B6": 6}
EPS = "1E-9"

# ---------------------------------------------------------------- styles (v19 look)
ARIAL = "Arial"
F_TITLE = Font(name=ARIAL, size=14, bold=True, color="FFFFFF")
F_HDR = Font(name=ARIAL, size=9, bold=True, color="FFFFFF")
F_GRP = Font(name=ARIAL, size=9, bold=True, color="1F4E79")
F_IN = Font(name=ARIAL, size=10, color="000000")
F_OUT = Font(name=ARIAL, size=10, bold=True, color="3F3F3F")
F_J = Font(name=ARIAL, size=12, bold=True, color="1F4E79")
F_TXT = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_HELP = Font(name=ARIAL, size=8, color="595959")
FILL = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731
P_BLUE, P_LIGHT, P_IN, P_OUT, P_J = FILL("2E75B6"), FILL("D9E1F2"), FILL("FFF2CC"), FILL("EDEDED"), FILL("EAF3DE")
P_GREEN_T, P_KEY, P_GREY, P_PURPLE = FILL("375623"), FILL("BDD7EE"), FILL("404040"), FILL("7030A0")
P_HELP = FILL("F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
LEFTW = Alignment(horizontal="left", vertical="top", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
KG = '0.000" kg"'
DATE = "mm/dd/yyyy"
RED, AMBER, GREEN = FILL("FFC7CE"), FILL("FFEB9C"), FILL("C6EFCE")
F_RED, F_AMBER, F_GREEN = Font(color="9C0006", bold=True), Font(color="9C5700", bold=True), Font(color="006100", bold=True)


def NORM(x):
    # UNICHAR(160) = NBSP everywhere (CHAR(160) depends on the code page: † on Mac Excel, not NBSP in LibreOffice)
    return 'UPPER(TRIM(CLEAN(SUBSTITUTE(%s,_xlfn.UNICHAR(160)," "))))' % x


def IS_CYC(x):
    return 'AND(LEN(%s)=2,LEFT(%s,1)="M",ISNUMBER(FIND(RIGHT(%s,1),"123456789")))' % (x, x, x)


def IS_LET(x):
    return 'AND(LEN(%s)=1,ISNUMBER(FIND(%s,"ABCDEFGH")))' % (x, x)


def ESC(x):  # escape COUNTIF/MATCH wildcards
    return 'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(%s,"~","~~"),"*","~*"),"?","~?")' % x


def SYM(f1, f2, t1, t2):
    """Slot-symmetric pick: if both slots flag, report the one with the smaller key."""
    return 'IF(AND(%s,%s),IF(@K1<=@K2,%s,%s),IF(%s,%s,IF(%s,%s,"")))' % (f1, f2, t1, t2, f1, t1, f2, t2)


BS = "'Bucket Status'!"
BS_A, BS_B, BS_C, BS_H, BS_N = (BS + "$%s$4:$%s$75" % (c, c) for c in "ABCHN")

# ---------------------------------------------------------------- machine-sheet column spec
# (name, header, template). Template tokens: @name = this row, ^name = previous row (initial value on
# the first data row), PREVLED = previous row's ledger range, HDR = ledger header range, {j} expanded first.
VISIBLE = [
    ("I", "Total drawn (kg)", '=IF(@uOK=1,@H+IF(@nCln="Y",Settings!$B$4+Settings!$B$5,0),"")'),
    ("J", "Powder on machine after build (kg)", '=IF(@uOK=1,@mL,"")'),
    ("K", "Draw from funnel", "=MID(" + "&".join(
        'IF(@tk%d>%s,"  +  "&TEXT(@tk%d,"0.000")&" kg "&@SK%d&" (funnel)","")' % (j, EPS, j, j) for j in range(1, 5))
     + ",6,500)"),
    ("L", "Draw from slot 1", '=IF(@d1>0,TEXT(@d1,"0.000")&" kg "&@K1,"")'),
    ("M", "Draw from slot 2", '=IF(@d2>0,TEXT(@d2,"0.000")&" kg "&@K2,"")'),
    ("N", "Slot 1 left", '=IF(AND(@act=1,@K1<>""),TEXT(IF(ABS(@lf1)<%s,0,@lf1),"0.000")&" kg "&@K1,"")' % EPS),
    ("O", "Slot 2 left", '=IF(AND(@act=1,@K2<>""),TEXT(IF(ABS(@lf2)<%s,0,@lf2),"0.000")&" kg "&@K2,"")' % EPS),
    ("Q", "LOT#s used (auto)", "=MID(" + "&".join('IF(@DK%d<>"",", "&@DL%d,"")' % (j, j) for j in range(1, 5))
     + '&IF(@f1=1,IF(@DK5<>"",", "&@DL5,"")&IF(@DK6<>"",", "&@DL6,""),IF(@DK6<>"",", "&@DL6,"")&IF(@DK5<>"",", "&@DL5,"")),3,500)'),
    ("R", "Row check",
     '=IF(@eMsg<>"","ERROR: "&@eMsg,IF(@wMsg<>"","WARN: "&@wMsg,IF(@mPend<>"","PENDING",IF(@act=1,"OK",""))))'),
]

H = []   # helper columns, in order


def h(name, header, tpl):
    H.append((name, header, tpl))


# 1) normalized inputs
h("nBld", "nBuild", "=" + NORM("@A"))
h("nCln", "nClean", "=" + NORM("@C"))
h("nC1", "nC1", "=" + NORM("@D"))
h("nB1", "nB1", "=" + NORM("@E"))
h("nC2", "nC2", "=" + NORM("@F"))
h("nB2", "nB2", "=" + NORM("@G"))
h("act", "Active (Build #)", '=IF(@nBld<>"",1,0)')
h("uOK", "Usage valid", "=IF(AND(@act=1,ISNUMBER(@H)),1,0)")
h("aRow", "Active row #", "=IF(@act=1,ROW(),0)")
h("LBD", "Last build date", "=IF(AND(@act=1,ISNUMBER(@B)),@B,^LBD)")
# 2) slot state
for s in "12":
    h("EC" + s, "EffC%s" % s, '=IF(@nC{s}="EMPTY","",IF(@nC{s}<>"",@nC{s},^EC{s}))'.replace("{s}", s))
    h("EB" + s, "EffB%s" % s, '=IF(@nC{s}="EMPTY","",IF(@nB{s}<>"",@nB{s},^EB{s}))'.replace("{s}", s))
    h("K" + s, "Key%s" % s, '=IF(AND(@EC{s}<>"",@EB{s}<>""),@EC{s}&@EB{s}&"-","")'.replace("{s}", s))
    h("Ix" + s, "Ledger idx%s" % s,
      '=IF(@K{s}="",0,IFERROR(IF(EXACT(INDEX(HDR,1,MATCH(@K{s},HDR,0)),@K{s}),MATCH(@K{s},HDR,0),0),0))'.replace("{s}", s))
    h("Ld" + s, "Loaded%s (row)" % s, '=IF(@K{s}="","",IF(@K{s}=^K{s},^Ld{s},ROW()))'.replace("{s}", s))
    h("Rm" + s, "Removed%s" % s, '=IF(AND(^K{s}<>"",@K{s}<>^K{s}),1,0)'.replace("{s}", s))
    h("New" + s, "Loaded new%s" % s, '=IF(AND(@K{s}<>"",@K{s}<>^K{s}),1,0)'.replace("{s}", s))
# 3) removal order (longest-loaded first, tie by key) -> A then B
h("R1st", "Removed first (slot)",
  "=IF(AND(@Rm1=1,@Rm2=1),IF(OR(^Ld1<^Ld2,AND(^Ld1=^Ld2,^K1<=^K2)),1,2),IF(@Rm1=1,1,IF(@Rm2=1,2,0)))")
h("RA", "Out A key", '=IF(@R1st=1,^K1,IF(@R1st=2,^K2,""))')
h("RAi", "Out A idx", "=IF(@R1st=1,^Ix1,IF(@R1st=2,^Ix2,0))")
h("RB", "Out B key", '=IF(AND(@Rm1=1,@Rm2=1),IF(@R1st=1,^K2,^K1),"")')
h("RBi", "Out B idx", "=IF(AND(@Rm1=1,@Rm2=1),IF(@R1st=1,^Ix2,^Ix1),0)")
h("wA", "Out A leftover", '=IF(@RA="",0,MAX(0,IF(@RAi=0,0,INDEX(PREVLED,1,@RAi))))')
h("wB", "Out B leftover", '=IF(@RB="",0,MAX(0,IF(@RBi=0,0,INDEX(PREVLED,1,@RBi))))')
# 4) funnel queue: start (prev end) -> append -> consume -> compact
h("n0", "Funnel lots at start", '=(^FK1<>"")+(^FK2<>"")+(^FK3<>"")+(^FK4<>"")')
h("s0", "Funnel kg at start", "=^FQ1+^FQ2+^FQ3+^FQ4")
h("pA", "Into funnel A", "=MIN(@wA,MAX(0,Settings!$B$6-@s0))")
h("aA", "A appended", "=IF(@pA>%s,1,0)" % EPS)
h("pB", "Into funnel B", "=MIN(@wB,MAX(0,Settings!$B$6-(@s0+IF(@aA=1,@pA,0))))")
h("aB", "B appended", "=IF(@pB>%s,1,0)" % EPS)
h("mLots", "Funnel lots after append", "=@n0+@aA+@aB")
for j in range(1, 5):
    h("SK%d" % j, "Funnel lot %d key (start)" % j,
      '=IF({j}<=@n0,^FK{j},IF(AND(@aA=1,{j}=@n0+1),@RA,IF(AND(@aB=1,{j}=@n0+@aA+1),@RB,"")))'.replace("{j}", str(j)))
for j in range(1, 5):
    h("SQ%d" % j, "Funnel lot %d kg (start)" % j,
      "=IF({j}<=@n0,^FQ{j},IF(AND(@aA=1,{j}=@n0+1),@pA,IF(AND(@aB=1,{j}=@n0+@aA+1),@pB,0)))".replace("{j}", str(j)))
h("nd0", "Need (total)", "=IF(@uOK=1,@I,0)")
for j in range(1, 5):
    h("tk%d" % j, "Funnel lot %d used" % j,
      "=IF(MIN(@SQ{j},@nd{i})>0,MIN(@SQ{j},@nd{i}),0)".replace("{j}", str(j)).replace("{i}", str(j - 1)))
    h("nd%d" % j, "Need after lot %d" % j, "=@nd{i}-@tk{j}".replace("{j}", str(j)).replace("{i}", str(j - 1)))
for j in range(1, 5):
    h("rq%d" % j, "Funnel lot %d left" % j, "=@SQ{j}-@tk{j}".replace("{j}", str(j)))
h("drop", "Lots emptied", "=" + "+".join('AND(@SK%d<>"",@rq%d<=%s)' % (j, j, EPS) for j in range(1, 5)))
for j in range(1, 5):
    h("FK%d" % j, "Funnel lot %d key (end)" % j,
      '=IF({j}+@drop<=MIN(@mLots,4),INDEX(@SK1:@SK4,1,{j}+@drop),"")'.replace("{j}", str(j)))
for j in range(1, 5):
    h("FQ%d" % j, "Funnel lot %d kg (end)" % j,
      "=IF({j}+@drop<=MIN(@mLots,4),INDEX(@rq1:@rq4,1,{j}+@drop),0)".replace("{j}", str(j)))
h("fAft", "Funnel kg (end)", "=@FQ1+@FQ2+@FQ3+@FQ4")
h("nS", "Need after funnel", "=IF(@nd4<%s,0,@nd4)" % EPS)
# 5) slot draws
h("av1", "Slot1 avail", '=IF(@K1="","",IF(@Ix1=0,0,INDEX(PREVLED,1,@Ix1)))')
h("av2", "Slot2 avail", '=IF(@K2="","",IF(@Ix2=0,0,INDEX(PREVLED,1,@Ix2)))')
h("on", "Buckets loaded", '=(@K1<>"")+(@K2<>"")')
h("f1", "Slot1 drains first", '=IF(@on=2,IF(OR(@Ld1<@Ld2,AND(@Ld1=@Ld2,@K1<=@K2)),1,0),IF(@K1<>"",1,0))')
h("tF", "First bucket draw", "=IF(@uOK=0,0,IF(@on=2,MIN(MAX(0,IF(@f1=1,@av1,@av2)),@nS),IF(@on=1,@nS,0)))")
h("nR", "Need after first", "=IF(@nS-@tF<%s,0,@nS-@tF)" % EPS)
h("tS", "Second bucket draw", "=IF(AND(@uOK=1,@on=2),@nR,0)")
h("su1", "Slot1 used", '=IF(@K1="",0,IF(@f1=1,@tF,@tS))')
h("su2", "Slot2 used", '=IF(@K2="",0,IF(@f1=1,@tS,@tF))')
h("noB", "Unsupplied kg", "=IF(AND(@uOK=1,@on=0,@nS>%s),@nS,0)" % EPS)
h("d1", "Slot1 draw", "=IF(@su1>%s,@su1,0)" % EPS)
h("d2", "Slot2 draw", "=IF(@su2>%s,@su2,0)" % EPS)
h("lf1", "Slot1 left", '=IF(@K1="","",@av1-@su1)')
h("lf2", "Slot2 left", '=IF(@K2="","",@av2-@su2)')
h("mL", "Powder on machine", "=N(@lf1)+N(@lf2)+@fAft")
h("lastS", "Last-drawn slot", '=IF(@on=2,IF(@f1=1,2,1),IF(@K1<>"",1,IF(@K2<>"",2,0)))')
h("ovd", "Overdraw", "=IF(AND(@uOK=1,@lastS>0),IF(IF(@lastS=1,@lf1,@lf2)<-%s,1,0),0)" % EPS)
# 6) draw columns (key block then kg block, contiguous) + LOT#s
for j in range(1, 5):
    h("DK%d" % j, "Draw %d key (funnel)" % j, '=IF(@tk{j}>0,@SK{j},"")'.replace("{j}", str(j)))
h("DK5", "Draw 5 key (slot 1)", '=IF(@d1>0,@K1,"")')
h("DK6", "Draw 6 key (slot 2)", '=IF(@d2>0,@K2,"")')
for j in range(1, 5):
    h("DQ%d" % j, "Draw %d kg" % j, "=@tk%d" % j)
h("DQ5", "Draw 5 kg", "=@d1")
h("DQ6", "Draw 6 kg", "=@d2")
for j in range(1, 7):
    h("DL%d" % j, "Draw %d LOT#" % j,
      ('=IF(@DK{j}="","",IFERROR(IF(INDEX(%s,MATCH(@DK{j},%s,0))="",@DK{j}&" (no LOT#)",'
       'INDEX(%s,MATCH(@DK{j},%s,0))),@DK{j}&" (not in Bucket Status)"))' % (BS_H, BS_A, BS_H, BS_A)).replace("{j}", str(j)))
# 7) ledger (inserted here; generated separately) ...
LEDGER_AFTER = "DL6"
# 8) per-slot flags and messages
for s in "12":
    h("unk" + s, "Not-filled flag%s" % s,
      "=IF(@New{s}=1,IF(@Ix{s}=0,1,IF(N(INDEX(%s,@Ix{s}))<=0,1,0)),0)".replace("{s}", s) % BS_B)
    h("dte" + s, "Fill-date flag%s" % s,
      ("=IF(AND(@New{s}=1,@unk{s}=0),IF(NOT(ISNUMBER(INDEX(%s,@Ix{s}))),1,IF(AND(ISNUMBER(@B),@B<INDEX(%s,@Ix{s})),2,0)),0)"
       % (BS_C, BS_C)).replace("{s}", s))
    h("both" + s, "Both-machines flag%s" % s,
      "=IF(AND(@New{s}=1,@Ix{s}>0),N(INDEX(%s,@Ix{s})),0)".replace("{s}", s) % BS_N)
    h("prt" + s, "Partial-entry flag%s" % s,
      ('=IF(AND(@K{s}<>"",OR(AND(@nB{s}<>"",@nC{s}=""),AND(@nC{s}<>"",@nC{s}<>"EMPTY",@nB{s}=""))),1,0)').replace("{s}", s))
    h("mFun" + s, "Oracle: in-funnel msg%s" % s,
      ('=IF(AND(@New{s}=1,OR(@SK1=@K{s},@SK2=@K{s},@SK3=@K{s},@SK4=@K{s})),'
       '@K{s}&" loaded while its leftover is still in the funnel","")').replace("{s}", s))
ws_any = ",".join(['AND(@A<>"",@nBld="")', 'AND(@B<>"",%s="")' % NORM("@B"), 'AND(@C<>"",@nCln="")',
                   'AND(@D<>"",@nC1="")', 'AND(@E<>"",@nB1="")', 'AND(@F<>"",@nC2="")', 'AND(@G<>"",@nB2="")',
                   'AND(@H<>"",%s="")' % NORM("@H")])
# --- ERROR message columns, in R priority order
ERR = [
    ("mWs", "Whitespace-only / non-printing input",
     '=IF(OR(%s),"Whitespace-only or non-printing entry - clear the cell (use EMPTY to empty a slot)","")' % ws_any,
     "Clear the cell (Delete key). To empty a slot type EMPTY in the Cycle column."),
    ("mCyc", "Invalid Cycle / Bucket entry",
     '=IF(OR(AND(@nC1<>"",@nC1<>"EMPTY",NOT(%s)),AND(@nC2<>"",@nC2<>"EMPTY",NOT(%s)),AND(@nB1<>"",NOT(%s)),'
     'AND(@nB2<>"",NOT(%s))),"Invalid entry: Cycle must be M1-M9 or EMPTY, Bucket must be A-H","")'
     % (IS_CYC("@nC1"), IS_CYC("@nC2"), IS_LET("@nB1"), IS_LET("@nB2")),
     "Correct the Cycle (M1-M9 or EMPTY) or Bucket (A-H) entry."),
    ("mEmp", "EMPTY entered with a bucket letter",
     '=IF(OR(AND(@nC1="EMPTY",@nB1<>""),AND(@nC2="EMPTY",@nB2<>"")),"EMPTY entered together with a bucket letter - clear the letter","")',
     "Clear the Bucket letter next to EMPTY."),
    ("mNoKey", "Bucket key can't form",
     '=IF(OR(AND(OR(@nC1<>"",@nB1<>""),@nC1<>"EMPTY",@K1=""),AND(OR(@nC2<>"",@nB2<>""),@nC2<>"EMPTY",@K2="")),'
     '"Key can\'t form - enter both Cycle and Bucket for the slot","")',
     "Enter both Cycle and Bucket the first time a bucket goes into a slot."),
    ("mNoBld", "Slot/usage entries on a row with no Build #",
     '=IF(AND(@act=0,OR(@C<>"",@D<>"",@E<>"",@F<>"",@G<>"",@H<>"")),"Entries on a row with no Build # - enter the Build # or clear the row","")',
     "Enter the Build # on that row, or clear the row."),
    ("mUse", "Usage non-numeric or negative",
     '=IF(AND(%s<>"",OR(NOT(ISNUMBER(@H)),N(@H)<0)),"Usage must be a number >= 0 (text or negative entered)","")' % NORM("@H"),
     "Re-type Calc. usage as a plain number (no text, no units), >= 0."),
    ("mCln", "Cleaned Chamber not Y/N/blank",
     '=IF(AND(@nCln<>"",@nCln<>"Y",@nCln<>"N"),"Cleaned Chamber must be Y, N or blank","")',
     "Enter Y, N or leave blank."),
    ("mSame", "Same bucket in both slots",
     '=IF(AND(@K1<>"",@K1=@K2),"Same bucket "&@K1&" in both slots","")',
     "A bucket can only sit in one slot. Correct the slot entries."),
    ("mUnk", "Bucket not in Bucket Status / Not filled",
     "=" + SYM("@unk1=1", "@unk2=1", '@K1&" is not in Bucket Status or is Not filled"', '@K2&" is not in Bucket Status or is Not filled"'),
     "Fill in Full weight and Date filled on Bucket Status before loading the bucket."),
    ("mDte", "Bucket loaded before its Date filled (or no Date filled)",
     "=" + SYM("@dte1>0", "@dte2>0",
               'IF(@dte1=1,@K1&" has no valid Date filled in Bucket Status (no LOT#)",@K1&" loaded before its Date filled")',
               'IF(@dte2=1,@K2&" has no valid Date filled in Bucket Status (no LOT#)",@K2&" loaded before its Date filled")'),
     "Check the build Date and the bucket's Date filled on Bucket Status."),
    ("mBoth", "Bucket used on both machines",
     "=" + SYM("@both1=1", "@both2=1", '@K1&" is used on both machines"', '@K2&" is used on both machines"'),
     "A bucket may only be used on one machine. Correct the bucket entries on one sheet."),
    ("mLot5", "More than 4 funnel lots",
     '=IF(@mLots>4,"More than 4 bucket leftovers in funnel","")',
     "Funnel holds at most 4 leftovers. Check the swap entries; run builds before further swaps."),
    ("mFunS", "Bucket loaded while its leftover is in the funnel",
     '=IF(AND(@mFun1<>"",@mFun2<>""),IF(@K1<=@K2,@mFun1,@mFun2),@mFun1&@mFun2)',
     "A removed bucket's leftover is still in the funnel - don't reload it until the funnel lot is used up."),
    ("mNoB", "Usage with no bucket and no funnel powder",
     '=IF(@noB>%s,TEXT(@noB,"0.000")&" kg used with no bucket loaded","")' % EPS,
     "Enter the bucket(s) on the machine for this build."),
    ("mOvd", "Overdraw (bucket goes negative)",
     '=IF(@ovd=1,"Overdraw: "&IF(@lastS=1,@K1,@K2)&" goes negative","")',
     "Usage exceeds the powder available. Check usage, bucket weights and swap entries."),
    ("mLed", "Ledger negative (any bucket < -0.0005 kg)",
     '=IF(AND(@ledMin<-0.0005,^ledMin>=-0.0005),"Ledger negative: a bucket\'s remaining went below 0","")',
     "See the Overdraw check; a bucket was drawn below zero."),
]
WARN = [
    ("mFFA", "Funnel full (overflow kept on bucket)",
     '=IF(@wA-@pA>%s,"Funnel full: "&TEXT(@wA-@pA,"0.000")&" kg of "&@RA&" did not fit (stays on bucket)","")' % EPS,
     "Informational: the part that did not fit stays on the bucket record."),
    ("mFFB", "Funnel full (overflow kept on bucket), 2nd bucket",
     '=IF(@wB-@pB>%s,"Funnel full: "&TEXT(@wB-@pB,"0.000")&" kg of "&@RB&" did not fit (stays on bucket)","")' % EPS,
     "Informational: the part that did not fit stays on the bucket record."),
    ("mDt", "Build date missing / not a date / out of order",
     '=IF(@act=0,"",IF(@B="","Build # has no Date",IF(NOT(ISNUMBER(@B)),"Date is not a real date",'
     'IF(@B<^LBD,"Date is earlier than the previous build\'s date",""))))',
     "Enter the build Date as mm/dd/yyyy (or confirm the order is intended)."),
    ("mDup", "Duplicate Build # on a sheet", None,  # generated (needs a growing range)
     "Confirm the repeat is intended (e.g. re-run); otherwise correct the Build #."),
    ("mPart", "Cycle-only or bucket-letter-only entry",
     '=IF(OR(@prt1=1,@prt2=1),"Cycle or bucket letter entered alone - resulting key: "&IF(AND(@prt1=1,@prt2=1),'
     'IF(@K1<=@K2,@K1&", "&@K2,@K2&", "&@K1),IF(@prt1=1,@K1,@K2)),"")',
     "Confirm the resulting key is the bucket actually loaded."),
    ("mPend", "Pending build (Build # but no usage)",
     '=IF(AND(@act=1,@H=""),"Pending: usage not entered","")',
     "Enter Calc. usage when the build is finished."),
]
for name, _, tpl, _ in ERR + WARN:
    if tpl is not None:
        h(name, name, tpl)
    else:
        h(name, name, "DUP")
h("ledMin", "Ledger min", "=MIN(LEDROW)")
h("eMsg", "First error", "=" + "".join('IF(@%s<>"",@%s,' % (n, n) for n, _, _, _ in ERR + [("mInt",)*4]) + '""' + ")" * (len(ERR) + 1))
h("wMsg", "First warning", "=" + "".join('IF(@%s<>"",@%s,' % (n, n) for n, _, _, _ in WARN if n != "mPend") + '""' + ")" * (len(WARN) - 1))
h("intB", "Integrity (part 2)", "INT")
h("mInt", "Formula integrity", "INT")
h("intR", "Integrity of R/eMsg", "INT")   # not read by R (would be circular); counted on Checks
INT_CHECK = ("mInt", "Formula overwritten on a machine-sheet row",
             "Formula overwritten on this row - restore the row's formulas from the template",
             "Restore the sheet from the template / previous version; never type over grey cells.")

INIT = {"LBD": "0", "ledMin": "0"}
for s in "12":
    INIT.update({"EC" + s: '""', "EB" + s: '""', "K" + s: '""', "Ld" + s: '""', "Ix" + s: "0"})
for j in range(1, 5):
    INIT.update({"FK%d" % j: '""', "FQ%d" % j: "0"})


class Layout:
    """Column letters for one machine sheet."""

    def __init__(self):
        self.col = {c: c for c in "ABCDEFGHIJKLMNOPQR"}
        self.helpers = []          # (name, header, tpl) in sheet order, ledger included as ("LED", key, idx)
        n = 19                     # S
        for name, header, tpl in H:
            self.col[name] = CL(n)
            self.helpers.append((name, header, tpl))
            n += 1
            if name == LEDGER_AFTER:
                self.led_first = n
                for k in range(72):
                    self.helpers.append(("LED%d" % k, None, None))
                    self.col["LED%d" % k] = CL(n)
                    n += 1
                self.led_last = n - 1
        self.last_col = n - 1
        self.hdr = "$%s$4:$%s$4" % (CL(self.led_first), CL(self.led_last))

    def led_range(self, r):
        return "$%s%d:$%s%d" % (CL(self.led_first), r, CL(self.led_last), r)


L = Layout()
TOKEN = re.compile(r"([@^])([A-Za-z_][A-Za-z0-9_]*)")


def expand(tpl, r, first=FIRST):
    def sub(m):
        kind, name = m.group(1), m.group(2)
        if kind == "@":
            return "%s%d" % (L.col[name], r)
        if r > first:
            return "%s%d" % (L.col[name], r - 1)
        return INIT[name]
    out = tpl.replace("PREVLED", L.led_range(r - 1 if r > first else 3)).replace("HDR", L.hdr)
    out = out.replace("LEDROW", L.led_range(r))
    return TOKEN.sub(sub, out)


def formula_columns():
    """Every formula column on a data row (for the integrity check)."""
    skip = ("mInt", "intB", "intR", "eMsg")      # R -> eMsg -> mInt: checking these from mInt would be circular
    cols = [c for c, _, _ in VISIBLE if c != "R"] + [L.col[n] for n, _, _ in L.helpers if n not in skip]
    return cols


@functools.lru_cache(maxsize=None)
def row_formulas(r, last):
    """{column letter: formula} for data row r (cached: identical for every workbook of the same size)."""
    f = {}
    for col, _, tpl in VISIBLE:
        f[col] = expand(tpl, r)
    for name, _, tpl in L.helpers:
        c = L.col[name]
        if name.startswith("LED"):
            prev = "%s%d" % (c, r - 1) if r > FIRST else "%s$3" % c
            f[c] = "=%s-SUMIF(%s%d:%s%d,%s$4,%s%d:%s%d)" % (
                prev, L.col["DK1"], r, L.col["DK6"], r, c, L.col["DQ1"], r, L.col["DQ6"], r)
        elif tpl == "DUP":
            nb = L.col["nBld"]
            if r == FIRST:
                f[c] = '=""'
            else:
                f[c] = ('=IF(%s%d=0,"",IF(COUNTIF($%s$%d:%s%d,"="&%s)>0,"Duplicate Build # on this sheet",""))'
                        % (L.col["act"], r, nb, FIRST, nb, r - 1, ESC("%s%d" % (nb, r))))
        elif tpl == "INT":
            pass
        else:
            f[c] = expand(tpl, r)
    cols = formula_columns()
    half = len(cols) // 2
    a, b = cols[:half], cols[half:]
    f[L.col["intB"]] = "=IF(AND(%s),1,0)" % ",".join("_xlfn.ISFORMULA(%s%d)" % (c, r) for c in b)
    f[L.col["mInt"]] = '=IF(AND(%s,%s%d=1),"","%s")' % (
        ",".join("_xlfn.ISFORMULA(%s%d)" % (c, r) for c in a), L.col["intB"], r, INT_CHECK[2])
    f[L.col["intR"]] = '=IF(AND(%s),"","%s")' % (
        ",".join("_xlfn.ISFORMULA(%s%d)" % (c, r) for c in ("R", L.col["eMsg"], L.col["mInt"], L.col["intB"])), INT_CHECK[2])
    return f


# ---------------------------------------------------------------- sheet builders
def _title(ws, text, ncols, fill=P_BLUE, size=14, height=30):
    ws.cell(1, 1, text).font = Font(name=ARIAL, size=size, bold=True, color="FFFFFF")
    for c in range(1, ncols + 1):
        ws.cell(1, c).fill = fill
    ws.cell(1, 1).alignment = CENTER
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = height


def _hdr(cell, text, fill=P_BLUE, font=F_HDR):
    cell.value = text
    cell.font, cell.fill, cell.alignment, cell.border = font, fill, CENTER, BORDER


def build_machine(ws, name, nrows, rows):
    last = FIRST + nrows - 1
    _title(ws, "EOS P110 — %s Build Log" % name, 18)
    groups = [("A2:C2", "Build info"), ("D2:E2", "Slot 1 bucket"), ("F2:G2", "Slot 2 bucket"), ("H2:H2", "Usage"),
              ("I2:J2", "Result"), ("K2:M2", "Powder source (funnel FIFO → longest-loaded bucket → other)"),
              ("N2:O2", "Slot remaining"), ("P2:P2", "Notes"), ("Q2:R2", "Traceability / check")]
    for rng, text in groups:
        a, b = rng.split(":")
        _hdr(ws[a], text, P_LIGHT, F_GRP)
        if a != b:
            ws.merge_cells(rng)
    hcol0, hcol1 = CL(19), CL(L.last_col)
    _hdr(ws["S2"], "Calculated detail (click + to expand)", P_LIGHT, F_GRP)
    ws.merge_cells("S2:%s2" % hcol1)
    # row 3 totals
    _hdr(ws["A3"], "Totals", P_BLUE)
    ws.merge_cells("A3:G3")
    for col, f in (("H", "=SUM(H%d:H%d)" % (FIRST, last)), ("I", "=SUM(I%d:I%d)" % (FIRST, last)),
                   ("K", "=SUM(%s%d:%s%d)" % (L.col["DQ1"], FIRST, L.col["DQ4"], last)),
                   ("L", "=SUM(%s%d:%s%d)" % (L.col["DQ5"], FIRST, L.col["DQ5"], last)),
                   ("M", "=SUM(%s%d:%s%d)" % (L.col["DQ6"], FIRST, L.col["DQ6"], last))):
        c = ws[col + "3"]
        c.value, c.number_format, c.font, c.fill, c.border, c.alignment = f, KG, F_OUT, P_OUT, BORDER, RIGHT
    # row 4 headers
    heads = {"A": "Build #", "B": "Date", "C": "Cleaned Chamber?", "D": "Slot 1 Cycle", "E": "Slot 1 Bucket",
             "F": "Slot 2 Cycle", "G": "Slot 2 Bucket", "H": "Calc. usage (kg)", "P": "Notes"}
    heads.update({c: t for c, t, _ in VISIBLE})
    for c, t in heads.items():
        _hdr(ws[c + "4"], t)
    for name_, header, _ in L.helpers:
        c = L.col[name_]
        cell = ws["%s4" % c]
        if name_.startswith("LED"):
            k = int(name_[3:])
            cell.value = "=%sA%d" % (BS, BS_FIRST + k)
            ws["%s3" % c] = "=N(%sB%d)" % (BS, BS_FIRST + k)
            ws["%s3" % c].number_format = "0.000"
            ws["%s3" % c].font = F_HELP
        else:
            cell.value = header
        cell.font, cell.fill, cell.alignment, cell.border = F_HDR, P_GREY, CENTER, BORDER
    ws.row_dimensions[4].height = 39.75
    # data rows
    numfmt = {"H": KG, "I": KG, "J": KG, "B": DATE}
    for r in range(FIRST, last + 1):
        for c, f in row_formulas(r, last).items():
            ws["%s%d" % (c, r)] = f
        for c in "ABCDEFGHP":
            cell = ws["%s%d" % (c, r)]
            cell.fill, cell.font, cell.border = P_IN, F_IN, BORDER
            cell.protection = Protection(locked=False)
            if c in numfmt:
                cell.number_format = numfmt[c]
        for c in "IKLMNOQR":
            cell = ws["%s%d" % (c, r)]
            cell.fill, cell.font, cell.border = P_OUT, F_OUT if c == "I" else F_TXT, BORDER
            if c == "I":
                cell.number_format = KG
        ws["J%d" % r].fill, ws["J%d" % r].font, ws["J%d" % r].number_format = P_J, F_J, KG
        ws["J%d" % r].border = BORDER
    for i, row in enumerate(rows or []):
        r = FIRST + i
        for c in "ABCDEFGHP":
            v = row.get(c)
            if v is not None:
                ws["%s%d" % (c, r)] = v
    # widths, freeze, grouping
    widths = {"A": 13, "B": 12, "C": 11, "D": 9, "E": 9, "F": 9, "G": 9, "H": 13, "I": 13, "J": 16, "K": 30,
              "L": 16, "M": 16, "N": 16, "O": 16, "P": 22, "Q": 34, "R": 44}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A5"
    ws.column_dimensions.group(hcol0, hcol1, hidden=True, outline_level=1)
    ws.sheet_properties.outlinePr.summaryRight = False
    # conditional formatting
    rng = lambda c: "%s%d:%s%d" % (c, FIRST, c, last)  # noqa: E731
    lc = lambda n: "$%s%d" % (L.col[n], FIRST)  # noqa: E731
    for c in ("R", "A"):
        ws.conditional_formatting.add(rng(c), FormulaRule(formula=['LEFT($R%d,5)="ERROR"' % FIRST], fill=RED, font=F_RED, stopIfTrue=True))
        ws.conditional_formatting.add(rng(c), FormulaRule(formula=['LEFT($R%d,4)="WARN"' % FIRST], fill=AMBER, font=F_AMBER, stopIfTrue=True))
    ws.conditional_formatting.add(rng("J"), FormulaRule(formula=["AND(ISNUMBER($J%d),$J%d<0)" % (FIRST, FIRST)], fill=RED, stopIfTrue=True))
    ws.conditional_formatting.add(rng("J"), FormulaRule(formula=["AND(ISNUMBER($J%d),$J%d<=Settings!$B$3)" % (FIRST, FIRST)], fill=AMBER, stopIfTrue=True))
    ws.conditional_formatting.add(rng("J"), FormulaRule(formula=["AND(ISNUMBER($J%d),$J%d>Settings!$B$3)" % (FIRST, FIRST)], fill=GREEN, stopIfTrue=True))
    bad = {
        "A": 'AND($A{r}<>"",%s="")' % lc("nBld"),
        "B": 'AND($B{r}<>"",NOT(ISNUMBER($B{r})))',
        "C": 'AND($C{r}<>"",%s<>"Y",%s<>"N")' % (lc("nCln"), lc("nCln")),
        "D": 'AND($D{r}<>"",%s<>"EMPTY",NOT(%s))' % (lc("nC1"), IS_CYC(lc("nC1"))),
        "F": 'AND($F{r}<>"",%s<>"EMPTY",NOT(%s))' % (lc("nC2"), IS_CYC(lc("nC2"))),
        "E": 'AND($E{r}<>"",OR(NOT(%s),%s="EMPTY"))' % (IS_LET(lc("nB1")), lc("nC1")),
        "G": 'AND($G{r}<>"",OR(NOT(%s),%s="EMPTY"))' % (IS_LET(lc("nB2")), lc("nC2")),
        "H": 'AND($H{r}<>"",OR(NOT(ISNUMBER($H{r})),N($H{r})<0))',
    }
    for c, f in bad.items():
        ws.conditional_formatting.add(rng(c), FormulaRule(formula=[f.replace("{r}", str(FIRST))], fill=RED, font=F_RED))
    # data validation (Stop style)
    def dv(sqref, **kw):
        d = DataValidation(showErrorMessage=True, showInputMessage=True, errorStyle="stop", allow_blank=True, **kw)
        d.add(sqref)
        ws.add_data_validation(d)
    dv(rng("B"), type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2099,12,31)",
       promptTitle="Date", prompt="Build date, mm/dd/yyyy", errorTitle="Date", error="Enter a date between 2020 and 2099.")
    dv(rng("C"), type="list", formula1="Lists!$C$2:$C$3", promptTitle="Cleaned chamber?",
       prompt="Y adds roller fill (Settings B4+B5). N or blank = no.", errorTitle="Cleaned chamber", error="Y, N or blank only.")
    for c in "DF":
        dv(rng(c), type="list", formula1="Lists!$A$2:$A$11", promptTitle="Cycle",
           prompt="M1-M9 when loading a bucket; EMPTY when the slot is emptied; blank = unchanged.",
           errorTitle="Cycle", error="M1-M9 or EMPTY only.")
    for c in "EG":
        dv(rng(c), type="list", formula1="Lists!$B$2:$B$9", promptTitle="Bucket",
           prompt="A-H when loading a bucket; blank = unchanged.", errorTitle="Bucket", error="A-H only.")
    dv(rng("H"), type="decimal", operator="between", formula1="0", formula2="50", promptTitle="Calc. usage (kg)",
       prompt="Powder used by this build in kg (0-50). Leave blank until known (row shows PENDING).",
       errorTitle="Usage", error="Number between 0 and 50 kg.")
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def build_bucket_status(ws, nrows, buckets):
    last = FIRST + nrows - 1
    _title(ws, "Bucket Status — 9 mix cycles × 8 buckets", 9, P_GREEN_T, 13)
    grp = [("A2", "Bucket key"), ("B2:C2", "Last fill"), ("D2", "Total drawn"), ("E2", "Est. remaining"),
           ("F2", "Status"), ("G2", "Notes"), ("H2", "LOT#"), ("I2", "Location")]
    for rng_, t in grp:
        _hdr(ws[rng_.split(":")[0]], t, P_LIGHT, F_GRP)
        if ":" in rng_:
            ws.merge_cells(rng_)
    heads = ["Bucket key", "Full weight (kg)", "Date filled", "Total drawn (kg)", "Est. remaining (kg)", "Status",
             "Notes", "LOT# (auto)", "Location (auto)",
             "Drawn Old (kg)", "Drawn New (kg)", "Loaded rows Old", "Loaded rows New", "On both machines",
             "Key sequence bad", "Formula overwritten", "Entry invalid"]
    for i, t in enumerate(heads):
        _hdr(ws.cell(3, i + 1), t, P_GREEN_T if i < 9 else P_GREY)
    O, N_ = "'Old P110'!", "'New P110'!"
    def sk(sheet, name, r):
        return "%s$%s$%d" % (sheet, L.col[name], r)
    for k, key in enumerate(KEYS):
        r = BS_FIRST + k
        lc = L.col["LED%d" % k]
        ws["A%d" % r] = key
        ws["D%d" % r] = "=J%d+K%d" % (r, r)
        ws["E%d" % r] = '=IF(B{r}="","Not filled",IF(ISNUMBER(B{r}),ROUND(B{r}-D{r},3),"INVALID"))'.replace("{r}", str(r))
        ws["F%d" % r] = ('=IF(B{r}="","Not filled",IF(NOT(ISNUMBER(E{r})),"INVALID",IF(E{r}<=0,"Empty",'
                         'IF(E{r}<Settings!$B$3,"Low","OK"))))').replace("{r}", str(r))
        ws["H%d" % r] = ('=IF(C{r}="","",IF(ISNUMBER(C{r}),A{r}&YEAR(C{r})&TEXT(MONTH(C{r}),"00")&TEXT(DAY(C{r}),"00"),'
                         '"INVALID DATE"))').replace("{r}", str(r))
        loc = '"Off machine"'
        parts = []
        for sh, nm in ((O, "Old P110"), (N_, "New P110")):
            parts.append(("A%d=%s" % (r, sk(sh, "K1", last)), '"%s slot 1"' % nm))
            parts.append(("A%d=%s" % (r, sk(sh, "K2", last)), '"%s slot 2"' % nm))
            parts.append(("OR(%s)" % ",".join("A%d=%s" % (r, sk(sh, "FK%d" % j, last)) for j in range(1, 5)),
                          '"%s funnel"' % nm))
        f = '=IF(B%d="","",' % r
        for cond, val in parts:
            f += "IF(%s,%s," % (cond, val)
        f += loc + ")" * (len(parts) + 1)
        ws["I%d" % r] = f
        ws["J%d" % r] = "=N(B%d)-%s$%s$%d" % (r, O, lc, last)
        ws["K%d" % r] = "=N(B%d)-%s$%s$%d" % (r, N_, lc, last)
        for col, sh in (("L", O), ("M", N_)):
            ws["%s%d" % (col, r)] = "=COUNTIF(%s$%s$%d:$%s$%d,A%d)+COUNTIF(%s$%s$%d:$%s$%d,A%d)" % (
                sh, L.col["K1"], FIRST, L.col["K1"], last, r, sh, L.col["K2"], FIRST, L.col["K2"], last, r)
        ws["N%d" % r] = "=IF(AND(L%d>0,M%d>0),1,0)" % (r, r)
        ws["O%d" % r] = "=IF(EXACT(A%d,Lists!D%d),0,1)" % (r, r - 2)
        ws["P%d" % r] = "=IF(AND(%s),0,1)" % ",".join("_xlfn.ISFORMULA(%s%d)" % (c, r) for c in "DEFHIJKLMNO")
        ws["Q%d" % r] = ('=IF(OR(AND(B{r}<>"",OR(NOT(ISNUMBER(B{r})),N(B{r})<=0)),AND(C{r}<>"",NOT(ISNUMBER(C{r})))),1,0)'
                         .replace("{r}", str(r)))
        b = (buckets or {}).get(key, {})
        for col in "BCG":
            if b.get(col) is not None:
                ws["%s%d" % (col, r)] = b[col]
        for col in "ABCDEFGHI":
            cell = ws["%s%d" % (col, r)]
            cell.border, cell.font = BORDER, F_TXT
        ws["A%d" % r].fill, ws["A%d" % r].font = P_KEY, F_BOLD
        for col in "BCG":
            ws["%s%d" % (col, r)].fill = P_IN
            ws["%s%d" % (col, r)].protection = Protection(locked=False)
        for col in "DEFHI":
            ws["%s%d" % (col, r)].fill = P_OUT
        ws["B%d" % r].number_format = KG
        ws["C%d" % r].number_format = DATE
        ws["D%d" % r].number_format = KG
        ws["E%d" % r].number_format = KG
        ws["F%d" % r].font = F_BOLD
        for col in "JK":
            ws["%s%d" % (col, r)].number_format = "0.000"
    # legend (v19 rows 77-81)
    _hdr(ws["A77"], "Status key", P_GREY)
    ws.merge_cells("A77:I77")
    for i, (t, fill) in enumerate([("OK — powder available", GREEN), ("Low — below warning threshold", AMBER),
                                   ("Empty — fully depleted", RED), ("Not filled — no weight recorded", P_OUT)]):
        r = 78 + i
        ws["A%d" % r].fill = fill
        ws["B%d" % r] = t
        ws["B%d" % r].font = F_TXT
        ws.merge_cells("B%d:I%d" % (r, r))
    rng_ = "F4:F75"
    for txt, fill, font in (("OK", GREEN, F_GREEN), ("Low", AMBER, F_AMBER), ("Empty", RED, F_RED)):
        ws.conditional_formatting.add(rng_, FormulaRule(formula=['$F4="%s"' % txt], fill=fill, font=font))
    ws.conditional_formatting.add("B4:C75", FormulaRule(formula=["$Q4=1"], fill=RED, font=F_RED))
    for c, w in zip("ABCDEFGHI", (10, 16, 14, 18, 20, 12, 22, 18, 20)):
        ws.column_dimensions[c].width = w
    ws.column_dimensions.group("J", "Q", hidden=True, outline_level=1)
    ws.freeze_panes = "B4"
    d = DataValidation(type="decimal", operator="between", formula1="0", formula2="20", showErrorMessage=True,
                       showInputMessage=True, errorStyle="stop", allow_blank=True, promptTitle="Full weight",
                       prompt="Bucket full weight in kg (0-20).", errorTitle="Full weight", error="Number 0-20 kg.")
    d.add("B4:B75")
    ws.add_data_validation(d)
    d = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2099,12,31)",
                       showErrorMessage=True, showInputMessage=True, errorStyle="stop", allow_blank=True,
                       promptTitle="Date filled", prompt="Date the bucket was filled, mm/dd/yyyy (sets the LOT#).",
                       errorTitle="Date filled", error="Enter a date 2020-2099.")
    d.add("C4:C75")
    ws.add_data_validation(d)
    ws.protection.sheet = True


def build_settings(ws, settings):
    _title(ws, "Settings", 4, P_GREY, 13)
    rows = [(3, "Low powder warning (kg)", "Bucket remaining below this shows Low status"),
            (4, "Left roller capacity (kg)", "Powder to fill left dosing roller from empty"),
            (5, "Right roller capacity (kg)", "Powder to fill right dosing roller from empty"),
            (6, "Funnel capacity (kg)", "Max powder the shared funnel holds. When a bucket is removed, its leftover "
                                        "goes into the funnel and is used before the buckets.")]
    for r, label, note in rows:
        ws["A%d" % r] = label
        ws["A%d" % r].font = F_BOLD
        c = ws["B%d" % r]
        c.value = (settings or {}).get("B%d" % r, SETTINGS_DEFAULTS["B%d" % r])
        c.number_format, c.fill, c.font, c.border = KG, P_IN, Font(name=ARIAL, size=10, color="0000FF"), BORDER
        c.protection = Protection(locked=False)
        ws["D%d" % r] = note
        ws["D%d" % r].font = F_TXT
        ws["D%d" % r].alignment = LEFTW
        d = DataValidation(type="decimal", operator="greaterThanOrEqual" if r == 3 else "greaterThan", formula1="0",
                           showErrorMessage=True, errorStyle="stop", errorTitle=label,
                           error="Must be a number %s 0." % (">=" if r == 3 else ">"))
        d.add("B%d" % r)
        ws.add_data_validation(d)
    ws.row_dimensions[6].height = 30
    _hdr(ws["A8"], "— Reference only (not used in calculations) —", P_GREY)
    ws.merge_cells("A8:D8")
    ws["A9"], ws["A10"] = "Edge waste constant (kg)", "Edge waste factor (kg/mm)"
    for r in (9, 10):
        ws["A%d" % r].font = F_BOLD
        ws["D%d" % r] = "FOR REFERENCE — handled in the powder calculator, not linked to any formula"
        ws["D%d" % r].font = F_HELP
    for c, w in zip("ABCD", (30, 16, 4, 64)):
        ws.column_dimensions[c].width = w
    ws.protection.sheet = True


def build_lists(ws):
    ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["F1"] = "Cycle", "Bucket", "Cleaned", "Bucket keys", "Bucket keys (row)"
    for i, v in enumerate(["M%d" % c for c in range(1, 10)] + ["EMPTY"]):
        ws.cell(2 + i, 1, v)
    for i, v in enumerate("ABCDEFGH"):
        ws.cell(2 + i, 2, v)
    ws["C2"], ws["C3"] = "Y", "N"
    for i, k in enumerate(KEYS):
        ws.cell(2 + i, 4, k)
        ws.cell(2, 6 + i, k)
    ws.sheet_state = "hidden"
    ws.protection.sheet = True


def build_dashboard(ws, nrows):
    last = FIRST + nrows - 1
    _title(ws, "Machine Dashboard — Live Powder Status", 6, P_PURPLE, 13)
    ws["A2"] = "Fully automatic — reads the last row with a Build # on each machine sheet. Nothing to enter here."
    ws["A2"].font = F_HELP
    ws.merge_cells("A2:F2")
    ws["A3"] = "=Checks!A3"
    ws["A3"].font = Font(name=ARIAL, size=12, bold=True)
    ws["A3"].alignment = CENTER
    ws.merge_cells("A3:F3")
    ws.conditional_formatting.add("A3:F3", FormulaRule(formula=['$A$3="ALL CHECKS PASS"'], fill=GREEN, font=F_GREEN, stopIfTrue=True))
    ws.conditional_formatting.add("A3:F3", FormulaRule(formula=['LEFT($A$3,2)="0 "'], fill=AMBER, font=F_AMBER, stopIfTrue=True))
    ws.conditional_formatting.add("A3:F3", FormulaRule(formula=['$A$3<>""'], fill=RED, font=F_RED, stopIfTrue=True))
    labels = ["Last build #", "Print date", "Powder used (this build)", "LOT#s used", "Slot 1 bucket + remaining",
              "Slot 2 bucket + remaining", "Funnel contents", "Powder on machine"]
    for base, sheet in ((5, "Old P110"), (15, "New P110")):
        sh = "'%s'!" % sheet
        _hdr(ws["A%d" % (base - 1)], sheet, P_BLUE, Font(name=ARIAL, size=14, bold=True, color="FFFFFF"))
        ws.merge_cells("A%d:F%d" % (base - 1, base - 1))
        hc = "$H$%d" % base
        ws[hc.replace("$", "")] = "=MAX(%s$%s$%d:$%s$%d)" % (sh, L.col["aRow"], FIRST, L.col["aRow"], last)
        IX = lambda col: "INDEX(%s$%s$1:$%s$%d,%s)" % (sh, col, col, last, hc)  # noqa: E731
        fk = [IX(L.col["FK%d" % j]) for j in range(1, 5)]
        fq = [IX(L.col["FQ%d" % j]) for j in range(1, 5)]
        lots = "&".join('IF(%s<>"",%s%s&" "&TEXT(%s,"0.000")&" kg","")' % (fk[j], '", "&' if j else "", fk[j], fq[j])
                        for j in range(4))
        fields = [
            '=IF(%s=0,"—",%s)' % (hc, IX("A")),
            '=IF(%s=0,"—",IF(ISNUMBER(%s),%s,"— (no date)"))' % (hc, IX("B"), IX("B")),
            '=IF(%s=0,"—",IF(ISNUMBER(%s),TEXT(%s,"0.000")&" kg","— (usage not entered)"))' % (hc, IX("I"), IX("I")),
            '=IF(%s=0,"—",IF(%s="","—",%s))' % (hc, IX("Q"), IX("Q")),
            '=IF(%s=0,"—",IF(%s="","— empty —",%s))' % (hc, IX("N"), IX("N")),
            '=IF(%s=0,"—",IF(%s="","— empty —",%s))' % (hc, IX("O"), IX("O")),
            '=IF(%s=0,"—",TEXT(%s,"0.000")&" kg"&IF(%s="",""," ("&%s&")"))' % (hc, IX(L.col["fAft"]), fk[0], lots),
            '=IF(%s=0,"—",TEXT(%s,"0.000")&" kg")' % (hc, IX(L.col["mL"])),
        ]
        for i, (lab, f) in enumerate(zip(labels, fields)):
            r = base + i
            ws["A%d" % r] = lab
            ws["A%d" % r].font, ws["A%d" % r].fill, ws["A%d" % r].border = F_BOLD, P_LIGHT, BORDER
            ws["B%d" % r] = f
            ws["B%d" % r].font = Font(name=ARIAL, size=11, bold=True)
            ws["B%d" % r].alignment = LEFT
            ws["B%d" % r].border = BORDER
            if i == 1:
                ws["B%d" % r].number_format = DATE
            ws.merge_cells("B%d:F%d" % (r, r))
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20
    for c in "CDEF":
        ws.column_dimensions[c].width = 14
    ws.column_dimensions["H"].hidden = True
    ws.protection.sheet = True


CHECK_ROW0 = 6


def checks_list(nrows):
    """(severity, label, count formula, first-occurrence formula, fix) for every check."""
    last = FIRST + nrows - 1
    out = []
    for sev, items in (("ERROR", ERR + [INT_CHECK[:2] + (None, INT_CHECK[3])]), ("WARN", WARN)):
        for name, label, _, fix in items:
            parts = []
            for col in [L.col[name]] + ([L.col["intR"]] if name == "mInt" else []):
                for sheet in MACHINES:
                    rng = "'%s'!$%s$%d:$%s$%d" % (sheet, col, FIRST, col, last)
                    parts.append((sheet, rng))
            cnt = "+".join('COUNTIF(%s,"?*")' % rng for _, rng in parts)
            first = '"—"'
            for sheet, rng in reversed(parts):
                first = ('IF(COUNTIF({rng},"?*")>0,"{sh} row "&(MATCH("?*",{rng},0)+{off})&": "&INDEX({rng},MATCH("?*",{rng},0)),{nxt})'
                         .format(rng=rng, sh=sheet, off=FIRST - 1, nxt=first))
            out.append((sev, label, "=" + cnt, "=" + first, fix))
    hdr_bad = "+".join("SUMPRODUCT(--NOT(EXACT('%s'!%s,Lists!$F$2:$%s$2)))" % (s, L.hdr, CL(5 + 72)) for s in MACHINES)
    hdr_int_parts = []
    for s in MACHINES:
        cells = ["_xlfn.ISFORMULA('%s'!%s%d)" % (s, CL(c), rr) for c in range(L.led_first, L.led_last + 1) for rr in (3, 4)]
        hdr_int_parts.append("IF(AND(%s),0,1)+IF(AND(%s),0,1)" % (",".join(cells[:72]), ",".join(cells[72:])))
    hdr_int = "+".join(hdr_int_parts)
    dash_cells = ["Machine Dashboard!A3"] + ["Machine Dashboard!B%d" % r for r in list(range(5, 13)) + list(range(15, 23))]
    dash_int = "IF(AND(%s),0,1)" % ",".join("_xlfn.ISFORMULA('%s'!%s)" % tuple(c.split("!")) for c in dash_cells)
    rowsused = "MAX(0,MAX(%s)-%d)" % (",".join("'%s'!$%s$%d:$%s$%d" % (s, L.col["aRow"], FIRST, L.col["aRow"], last)
                                       for s in MACHINES), FIRST - 1)
    sets = ("OR(NOT(ISNUMBER(Settings!B3)),NOT(ISNUMBER(Settings!B4)),NOT(ISNUMBER(Settings!B5)),NOT(ISNUMBER(Settings!B6)),"
            "N(Settings!B3)<0,N(Settings!B4)<=0,N(Settings!B5)<=0,N(Settings!B6)<=0)")
    first_bs = lambda col: ('IF(SUM(\'Bucket Status\'!${c}$4:${c}$75)>0,"Bucket Status row "&(MATCH(1,\'Bucket Status\'!${c}$4:${c}$75,0)+3)'  # noqa: E731
                            '&" ("&INDEX(\'Bucket Status\'!$A$4:$A$75,MATCH(1,\'Bucket Status\'!${c}$4:${c}$75,0))&")","—")').format(c=col)
    out += [
        ("ERROR", "Bucket Status key sequence not exactly M1A-…M9H-", "=SUM('Bucket Status'!$O$4:$O$75)",
         "=" + first_bs("O"), "Restore column A of Bucket Status (M1A- … M9H-, trailing dash)."),
        ("ERROR", "Bucket Status formula overwritten (D/E/F/H/I or hidden helpers)", "=SUM('Bucket Status'!$P$4:$P$75)",
         "=" + first_bs("P"), "Restore the sheet from the template; never type over grey cells."),
        ("ERROR", "Bucket Status entry invalid (weight not > 0, or date not a date)", "=SUM('Bucket Status'!$Q$4:$Q$75)",
         "=" + first_bs("Q"), "Re-type Full weight as a number and Date filled as mm/dd/yyyy."),
        ("ERROR", "Ledger headers don't match the 72 bucket keys", "=" + hdr_bad, '=IF(%s>0,"machine sheet ledger header row 4","—")' % hdr_bad,
         "Restore the machine sheet from the template."),
        ("ERROR", "Ledger header / opening-balance formulas overwritten", "=" + hdr_int, '=IF(%s>0,"machine sheet ledger rows 3-4","—")' % hdr_int,
         "Restore the machine sheet from the template."),
        ("ERROR", "Dashboard formula overwritten", "=" + dash_int, '=IF(%s>0,"Machine Dashboard","—")' % dash_int,
         "Restore the Machine Dashboard from the template."),
        ("ERROR", "Settings invalid (not a number, or <= 0)", "=IF(%s,1,0)" % sets, '=IF(%s,"Settings B3:B6","—")' % sets,
         "Enter numbers in Settings B3:B6 (B4-B6 > 0)."),
        ("WARN", "900 or more build rows used (sheet holds %d)" % nrows, "=IF(%s>=900,1,0)" % rowsused,
         '="Rows used: "&(%s)' % rowsused, "Start a new workbook (carry over loaded buckets and funnel) before row %d." % last),
        ("WARN", "Settings differ from defaults (1 / 0.6 / 0.6 / 6)",
         "=IF(OR(Settings!B3<>1,Settings!B4<>0.6,Settings!B5<>0.6,Settings!B6<>6),1,0)",
         '=IF(OR(Settings!B3<>1,Settings!B4<>0.6,Settings!B5<>0.6,Settings!B6<>6),"Settings B3:B6","—")',
         "Confirm the change is approved and recorded in the Change Log."),
    ]
    return out


def build_checks(ws, nrows):
    _title(ws, "Checks — the workbook's safety net (fix every ERROR; review every WARN)", 5, P_GREY, 13)
    ws["A2"] = ("Data validation can be bypassed by pasting, so this sheet re-checks every input. "
                "Row-level messages also appear in column R of each machine sheet.")
    ws["A2"].font = F_HELP
    ws.merge_cells("A2:E2")
    items = checks_list(nrows)
    n = len(items)
    r0, r1 = CHECK_ROW0, CHECK_ROW0 + n - 1
    errs = 'SUMIF($A$%d:$A$%d,"ERROR",$C$%d:$C$%d)' % (r0, r1, r0, r1)
    warns = 'SUMIF($A$%d:$A$%d,"WARN",$C$%d:$C$%d)' % (r0, r1, r0, r1)
    ws["A3"] = '=IF(%s+%s=0,"ALL CHECKS PASS",%s&" ERRORS / "&%s&" WARNINGS")' % (errs, warns, errs, warns)
    ws["A3"].font = Font(name=ARIAL, size=16, bold=True)
    ws["A3"].alignment = CENTER
    ws.merge_cells("A3:E3")
    ws.row_dimensions[3].height = 32
    ws.conditional_formatting.add("A3:E3", FormulaRule(formula=['$A$3="ALL CHECKS PASS"'], fill=GREEN, font=F_GREEN, stopIfTrue=True))
    ws.conditional_formatting.add("A3:E3", FormulaRule(formula=['LEFT($A$3,2)="0 "'], fill=AMBER, font=F_AMBER, stopIfTrue=True))
    ws.conditional_formatting.add("A3:E3", FormulaRule(formula=['$A$3<>""'], fill=RED, font=F_RED, stopIfTrue=True))
    for i, t in enumerate(["Severity", "Check", "Count", "First occurrence", "How to fix"]):
        _hdr(ws.cell(5, i + 1), t)
    for i, (sev, label, cnt, first, fix) in enumerate(items):
        r = r0 + i
        for c, v in zip("ABCDE", (sev, label, cnt, first, fix)):
            ws["%s%d" % (c, r)] = v
            ws["%s%d" % (c, r)].font = F_TXT
            ws["%s%d" % (c, r)].border = BORDER
            ws["%s%d" % (c, r)].alignment = LEFTW
        ws["C%d" % r].alignment = RIGHT
    rng = "A%d:E%d" % (r0, r1)
    ws.conditional_formatting.add(rng, FormulaRule(formula=['AND($A%d="ERROR",$C%d>0)' % (r0, r0)], fill=RED, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['AND($A%d="WARN",$C%d>0)' % (r0, r0)], fill=AMBER, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=["$C%d=0" % r0], fill=GREEN))
    for c, w in zip("ABCDE", (10, 52, 8, 70, 70)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A6"
    ws.protection.sheet = True


def _text_sheet(ws, title, blocks, widths):
    _title(ws, title, len(widths), P_GREY, 13)
    r = 3
    for block in blocks:
        if block[0] == "h":
            _hdr(ws.cell(r, 1), block[1], P_BLUE)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(widths))
            r += 1
        elif block[0] == "p":
            ws.cell(r, 1, block[1]).font = F_TXT
            ws.cell(r, 1).alignment = LEFTW
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(widths))
            ws.row_dimensions[r].height = max(15, 15 * (1 + len(block[1]) // 140))
            r += 1
        elif block[0] == "t":
            for i, t in enumerate(block[1]):
                _hdr(ws.cell(r, i + 1), t, P_LIGHT, F_GRP)
            r += 1
            for row in block[2]:
                for i, v in enumerate(row):
                    c = ws.cell(r, i + 1, v)
                    c.font, c.alignment, c.border = F_TXT, LEFTW, BORDER
                r += 1
        elif block[0] == "sp":
            r += 1
    for i, w in enumerate(widths):
        ws.column_dimensions[CL(i + 1)].width = w
    ws.protection.sheet = True


PASTE = [
    ("h", "Redoing a migration by hand (normally migrate_v19.py does this)"),
    ("p", "1. Paste VALUES ONLY: Ctrl+Alt+V → Values (Excel) / Ctrl+Shift+V → Values only (LibreOffice). Never paste formats or formulas."),
    ("p", "2. Machine sheets (Old P110, New P110): copy the source block A5:H<last> and paste at A5. Then copy column P (Notes) "
          "separately and paste at P5. Never paste A:P or whole rows: I:O, Q:R and the hidden columns are formulas, and "
          "sheet protection blocks it anyway."),
    ("p", "3. Bucket Status: paste the B:C block (Full weight, Date filled) at B4, then column G (Notes) at G4. Column A (keys) "
          "is fixed: M1A- … M9H-, 72 rows; do not paste over it."),
    ("p", "4. Settings: paste B3:B6 only."),
    ("p", "5. Whitespace: a cell that only contains spaces is an ERROR. To show that a slot was emptied, type EMPTY in the Cycle "
          "column and leave the Bucket column blank."),
    ("p", "6. Open the Checks sheet. Fix every ERROR and review every WARN until the banner shows ALL CHECKS PASS "
          "(or only reviewed WARNINGS remain). Record what you changed in the Change Log."),
    ("p", "Pasting replaces the cells' data validation in Excel, so dropdowns may stop working on pasted cells. The Checks "
          "sheet re-checks every input regardless."),
]


def build_workbook(data=None, nrows=1000, change_log=None):
    data = data or {}
    wb = Workbook()
    wb.remove(wb.active)
    names = ["Checks", "Old P110", "New P110", "Bucket Status", "Machine Dashboard", "Settings", "Change Log",
             "Paste Instructions", "Lists"]
    sh = {n: wb.create_sheet(n) for n in names}
    build_checks(sh["Checks"], nrows)
    for m in MACHINES:
        build_machine(sh[m], m, nrows, data.get("machines", {}).get(m))
    build_bucket_status(sh["Bucket Status"], nrows, data.get("buckets"))
    build_dashboard(sh["Machine Dashboard"], nrows)
    build_settings(sh["Settings"], data.get("settings"))
    _text_sheet(sh["Change Log"], "Change Log — P110 Powder Tracker",
                change_log or [("p", "No entries yet.")], (14, 30, 40, 40, 40))
    _text_sheet(sh["Paste Instructions"], "Paste Instructions", PASTE, (120,))
    build_lists(sh["Lists"])
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Checks"
    wb.calculation.fullCalcOnLoad = True
    return wb


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "P110_Powder_Tracker_v20_template.xlsx"
    build_workbook().save(out)
    print("wrote", out, "helper columns S..%s, ledger %s..%s" % (CL(L.last_col), CL(L.led_first), CL(L.led_last)))
