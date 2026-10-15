"""P110 Powder Tracker v20 — automated tests (spec §7).

Every test builds a workbook with build_tracker.py, recalculates it in LibreOffice headless (UNO,
calculateAll) and reads the computed values back. The oracle is ref_model.py.

    python3 tests/test_tracker.py [v19.xlsx] [--seed N] [--fuzz 500] [--sym 300]   (default: tests/fixtures/ copy)
        -> runs T1..T11, writes TEST_REPORT.md, exit code 0 only if every test passed.
    (pytest-compatible too: set P110_V19=<v19.xlsx> and run pytest tests/test_tracker.py)
"""
import datetime
import json
import os
import random
import re
import sys
import tempfile
import time
import traceback
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import openpyxl  # noqa: E402

import build_tracker as bt  # noqa: E402
import migrate_v19 as mg  # noqa: E402
from lo_calc import LO, read_all  # noqa: E402
from ref_model import EPS, norm, run_machine  # noqa: E402

TOL = 1e-9
CFG = {"v19": os.environ.get("P110_V19", os.path.join(HERE, "fixtures", "P110_Powder_Tracker_v19_09-14-2026.xlsx")), "seed": 20261015, "fuzz": 500, "sym": 300,
       "deliverable": os.path.join(ROOT, "P110_Powder_Tracker_v20.xlsx")}
REPORT = {}
TMP = tempfile.mkdtemp(prefix="p110_tests_")
_LO = None


def lo():
    global _LO
    if _LO is None:
        _LO = LO().__enter__()
    return _LO


def note(test, **kw):
    REPORT.setdefault(test, {}).update(kw)


# ------------------------------------------------------------------ workbook helpers
class Calc:
    """Recalculated values of one workbook."""

    def __init__(self, path, nrows):
        doc = lo().open(path)
        self.recalc_s = lo().last_recalc_s
        self.errors = lo().error_cells(doc)
        self.a = read_all(doc)
        doc.close(True)
        self.nrows = nrows
        self.last = bt.FIRST + nrows - 1

    def v(self, sheet, ref):
        m = re.match(r"([A-Z]+)(\d+)$", ref)
        col = 0
        for ch in m.group(1):
            col = col * 26 + ord(ch) - 64
        try:
            return self.a[sheet][int(m.group(2)) - 1][col - 1]
        except IndexError:
            return ""

    def m(self, sheet, name, r):
        return self.v(sheet, "%s%d" % (bt.L.col[name], r))

    def led(self, sheet, k, r):
        return self.v(sheet, "%s%d" % (bt.L.col["LED%d" % k], r))


def build_calc(data, nrows, name="wb", mutate=None):
    path = os.path.join(TMP, "%s_%d.xlsx" % (name, random.randint(0, 10 ** 9)))
    wb = bt.build_workbook(data, nrows)
    if mutate:
        mutate(wb)
    wb.save(path)
    c = Calc(path, nrows)
    os.remove(path)
    return c


def num(x):
    return x if isinstance(x, float) else None


def close(a, b, tol=TOL):
    if a is None or b is None or a == "" or b == "":
        return (a in (None, "")) and (b in (None, ""))
    return abs(a - b) <= tol


def oracle_rows(rows, nrows):
    rr = mg.oracle_rows(rows)
    blank = dict(build=None, date=None, clean=None, s1c=None, s1b=None, s2c=None, s2b=None, usage=None)
    return rr + [dict(blank) for _ in range(nrows - len(rr))]


def run_oracle(data, nrows):
    s = data.get("settings") or {}
    cap = s.get("B6", 6)
    roller = s.get("B4", 0.6) + s.get("B5", 0.6)
    full = mg.full_weights(data)
    return {m: run_machine(oracle_rows(data["machines"].get(m, []), nrows), full, cap, roller) for m in bt.MACHINES}


def loaded_order(rows):
    """loaded_at per slot, replicated from the oracle's rule (for column Q consumption order)."""
    cyc, bkt, key, at, res = {"1": "", "2": ""}, {"1": "", "2": ""}, {"1": "", "2": ""}, {"1": None, "2": None}, []
    for i, row in enumerate(rows):
        prev = dict(key)
        for s in "12":
            c, b = norm(row["s%sc" % s]), norm(row["s%sb" % s])
            if c == "EMPTY":
                cyc[s], bkt[s] = "", ""
            else:
                if c:
                    cyc[s] = c
                if b:
                    bkt[s] = b
            key[s] = cyc[s] + bkt[s] + "-" if cyc[s] and bkt[s] else ""
            if key[s] != prev[s]:
                at[s] = i if key[s] else None
        res.append(dict(at))
    return res


def final_funnel(o_last):
    lots = [list(x) for x in o_last["funnel_before"]]
    for k, a in o_last["funnel_used"]:
        for lot in lots:
            if lot[0] == k and lot[1] > 0:
                lot[1] -= a
                break
    return [lot for lot in lots if lot[1] > EPS]


def lot_of(data, key):
    return mg.lot_of(data, key)


def expected_R(i, o, issues_row, raw, extra_warns=(), extra_errs=()):
    errs = [t for _, kind, t in issues_row if kind == "ERROR"]
    if not errs and extra_errs:
        return "ERROR: " + extra_errs[0]
    warns = [t for _, kind, t in issues_row if kind == "WARN"]
    if errs:
        lot5 = [t for t in errs if t.startswith("More than")]
        fun = sorted(t for t in errs if "still in the funnel" in t)
        nob = [t for t in errs if "no bucket loaded" in t]
        ovd = [t for t in errs if t.startswith("Overdraw")]
        return "ERROR: " + (lot5 + fun[:1] + nob + ovd)[0]
    if warns:
        return "WARN: " + warns[0]
    if extra_warns:
        return "WARN: " + extra_warns[0]
    prt = []
    for s, k in (("1", o["k1"]), ("2", o["k2"])):
        c, b = norm(raw["s%sc" % s]), norm(raw["s%sb" % s])
        if k and ((b and not c) or (c and c != "EMPTY" and not b)):
            prt.append(k)
    if prt:
        return "WARN: Cycle or bucket letter entered alone - resulting key: " + ", ".join(sorted(prt))
    if o["active"]:
        u = raw["usage"]
        if not (isinstance(u, (int, float)) and not isinstance(u, bool)):
            return "PENDING"
        return "OK"
    return ""


ORACLE_MSG_COLS = ("mFFA", "mFFB", "mLot5", "mFun1", "mFun2", "mNoB", "mOvd")


def compare_machine(c, sheet, data, nrows, oracle, check_R=True, stop_at_lot5=True):
    """Compare one machine sheet with the oracle on every row. Returns (#rows compared, mismatches list)."""
    out, draws, issues = oracle
    rows = oracle_rows(data["machines"].get(sheet, []), nrows)
    full = mg.full_weights(data)
    order = loaded_order(rows)
    bad = []
    stop = nrows
    if stop_at_lot5:
        lot5 = [i for i, kind, t in issues if t.startswith("More than")]
        if lot5:
            stop = min(lot5)
    cum = {}
    seen, last_date, prev_min = set(), None, 0.0
    for i in range(stop):
        o, r = out[i], bt.FIRST + i
        row_draws = [(b, a, s) for ii, b, a, s in draws if ii == i]
        for b, a, _ in row_draws:
            cum[b] = cum.get(b, 0.0) + a

        def chk(label, got, exp, tol=TOL):
            if isinstance(exp, str) or isinstance(got, str) and exp is None:
                ok = (got if got is not None else "") == (exp if exp is not None else "")
            else:
                ok = close(num(got) if got != "" else None, exp, tol)
            if not ok:
                bad.append((sheet, r, label, got, exp))
        chk("I", c.v(sheet, "I%d" % r), o["total"])
        chk("J", c.v(sheet, "J%d" % r), o["left"])
        chk("K1", c.m(sheet, "K1", r), o["k1"])
        chk("K2", c.m(sheet, "K2", r), o["k2"])
        fb = o["funnel_before"]
        for j in range(4):
            chk("SK%d" % (j + 1), c.m(sheet, "SK%d" % (j + 1), r), fb[j][0] if j < len(fb) else "")
            chk("SQ%d" % (j + 1), c.m(sheet, "SQ%d" % (j + 1), r), fb[j][1] if j < len(fb) else 0.0)
        fu = [(k, a) for k, a in o["funnel_used"] if a > EPS]
        wu = [(c.m(sheet, "SK%d" % j, r), c.m(sheet, "tk%d" % j, r)) for j in range(1, 5)
              if num(c.m(sheet, "tk%d" % j, r)) and c.m(sheet, "tk%d" % j, r) > EPS]
        if len(fu) != len(wu) or any(a[0] != b[0] or abs(a[1] - b[1]) > TOL for a, b in zip(fu, wu)):
            bad.append((sheet, r, "funnel_used", wu, fu))
        chk("su1", c.m(sheet, "su1", r), o["s1u"] or 0.0)
        chk("su2", c.m(sheet, "su2", r), o["s2u"] or 0.0)
        chk("lf1", c.m(sheet, "lf1", r), o["slot1_left"])
        chk("lf2", c.m(sheet, "lf2", r), o["slot2_left"])
        chk("fAft", c.m(sheet, "fAft", r), o["funnel_after"])
        chk("mL", c.m(sheet, "mL", r), o["machine_left"])
        # draw columns
        src = {1: "funnel", 2: "funnel", 3: "funnel", 4: "funnel", 5: "slot1", 6: "slot2"}
        wd = sorted((c.m(sheet, "DK%d" % j, r), round(c.m(sheet, "DQ%d" % j, r), 9), src[j]) for j in range(1, 7)
                    if c.m(sheet, "DK%d" % j, r) != "" and c.m(sheet, "DQ%d" % j, r) > EPS)
        od = sorted((b, round(a, 9), s) for b, a, s in row_draws if a > EPS)
        if len(wd) != len(od) or any(x[0] != y[0] or x[2] != y[2] or abs(x[1] - y[1]) > 1e-8 for x, y in zip(wd, od)):
            bad.append((sheet, r, "draws", wd, od))
        # visible text columns
        def kgtxt(a):
            return "%.3f" % (round(a + 1e-12, 3) if a >= 0 else a)
        k_exp = "  +  ".join("%s kg %s (funnel)" % (kgtxt(a), k) for k, a in fu)
        chk("K", c.v(sheet, "K%d" % r), k_exp)
        for s, col in (("1", "L"), ("2", "M")):
            a = [x for b, x, so in row_draws if so == "slot" + s]
            chk(col, c.v(sheet, "%s%d" % (col, r)), ("%s kg %s" % (kgtxt(a[0]), o["k" + s])) if a else "")
        for s, col in (("1", "N"), ("2", "O")):
            lf = o["slot%s_left" % s]
            exp = ""
            if o["active"] and o["k" + s]:
                exp = "%s kg %s" % (kgtxt(0.0 if abs(lf) < 1e-9 else lf), o["k" + s])
            got = c.v(sheet, "%s%d" % (col, r))
            if got != exp and not (got.replace("-0.000", "0.000") == exp.replace("-0.000", "0.000")):
                bad.append((sheet, r, col, got, exp))
        # Q: funnel FIFO, then longest-loaded bucket, then the other
        fun_keys = [k for k, a in fu]
        sl = [(order[i][s], o["k" + s], s) for s in "12" if any(so == "slot" + s and a > EPS for b, a, so in row_draws)]
        sl.sort(key=lambda t: (t[0], t[1]))
        q_exp = ", ".join(lot_of(data, k) for k in fun_keys + [t[1] for t in sl])
        chk("Q", c.v(sheet, "Q%d" % r), q_exp)
        # ledger = oracle remaining for every key
        for k, key in enumerate(bt.KEYS):
            exp = full.get(key, 0.0) - cum.get(key, 0.0)
            got = c.led(sheet, k, r)
            if not close(got, exp, 1e-9):
                bad.append((sheet, r, "LED " + key, got, exp))
        # oracle messages appear identically
        iss = [x for x in issues if x[0] == i]
        msgs = {c.m(sheet, n, r) for n in ORACLE_MSG_COLS} - {""}
        for _, kind, t in iss:
            if t not in msgs:
                bad.append((sheet, r, "oracle msg missing", sorted(msgs), t))
        if check_R:
            extra = []
            if o["active"]:
                dt = rows[i]["date"]
                if dt is None:
                    extra.append("Build # has no Date")
                elif not isinstance(dt, datetime.datetime):
                    extra.append("Date is not a real date")
                elif last_date is not None and dt < last_date:
                    extra.append("Date is earlier than the previous build's date")
                if isinstance(dt, datetime.datetime):
                    last_date = dt
                if norm(rows[i]["build"]) in seen:
                    extra.append("Duplicate Build # on this sheet")
                seen.add(norm(rows[i]["build"]))
            cur_min = min(full.get(key, 0.0) - cum.get(key, 0.0) for key in bt.KEYS)
            led_err = ["Ledger negative: a bucket's remaining went below 0"] if cur_min < -0.0005 <= prev_min else []
            prev_min = cur_min
            chk("R", c.v(sheet, "R%d" % r), expected_R(i, o, iss, rows[i], extra, led_err))
    return stop, bad


def compare_bucket_status(c, data, nrows, oracles):
    bad = []
    drawn = {}
    for m in bt.MACHINES:
        for _, b, a, _ in oracles[m][1]:
            drawn[b] = drawn.get(b, 0.0) + a
    loc = {}
    for m in bt.MACHINES:
        o = oracles[m][0][-1]
        for s in "12":
            if o["k" + s]:
                loc.setdefault(o["k" + s], "%s slot %s" % (m, s))
        for k, _ in final_funnel(o):
            loc.setdefault(k, "%s funnel" % m)
    full = mg.full_weights(data)
    for k, key in enumerate(bt.KEYS):
        r = 4 + k
        d = drawn.get(key, 0.0)
        if not close(c.v("Bucket Status", "D%d" % r), d, 1e-9):
            bad.append(("BS", key, "D", c.v("Bucket Status", "D%d" % r), d))
        if key in full:
            e = round(full[key] - d, 3)
            st = "Empty" if e <= 0 else ("Low" if e < (data.get("settings") or {}).get("B3", 1) else "OK")
            lot = lot_of(data, key)
            l = loc.get(key, "Off machine")
        else:
            e, st, lot, l = "Not filled", "Not filled", "", ""
            if data["buckets"].get(key, {}).get("C"):
                lot = lot_of(data, key)
        got_e = c.v("Bucket Status", "E%d" % r)
        if (isinstance(e, str) and got_e != e) or (not isinstance(e, str) and not close(got_e, e, 1e-9)):
            bad.append(("BS", key, "E", got_e, e))
        for col, exp in (("F", st), ("H", lot), ("I", l)):
            if c.v("Bucket Status", "%s%d" % (col, r)) != exp:
                bad.append(("BS", key, col, c.v("Bucket Status", "%s%d" % (col, r)), exp))
    return bad


def dashboard_expected(data, nrows, oracles):
    exp = {}
    for base, m in ((5, "Old P110"), (15, "New P110")):
        out = oracles[m][0]
        rows = oracle_rows(data["machines"].get(m, []), nrows)
        act = [i for i, o in enumerate(out) if o["active"]]
        if not act:
            continue
        i = act[-1]
        o = out[i]
        fl = final_funnel(o) if o["total"] is not None else [list(x) for x in o["funnel_before"]]
        def kg(x):
            return "%.3f" % (0.0 if abs(x) < 1e-9 else x)
        exp[m] = {
            "Last build #": str(rows[i]["build"]),
            "Powder used (this build)": (kg(o["total"]) + " kg") if o["total"] is not None else "— (usage not entered)",
            "Slot 1 bucket + remaining": ("%s kg %s" % (kg(o["slot1_left"]), o["k1"])) if o["k1"] else "— empty —",
            "Slot 2 bucket + remaining": ("%s kg %s" % (kg(o["slot2_left"]), o["k2"])) if o["k2"] else "— empty —",
            "Funnel contents": kg(o["funnel_after"]) + " kg" + (
                (" (" + ", ".join("%s %s kg" % (k, kg(q)) for k, q in fl) + ")") if fl else ""),
            "Powder on machine": kg(o["machine_left"]) + " kg",
        }
    return exp


def read_dashboard(c):
    d = {}
    for base, m in ((5, "Old P110"), (15, "New P110")):
        d[m] = {c.v("Machine Dashboard", "A%d" % (base + i)): c.v("Machine Dashboard", "B%d" % (base + i)) for i in range(8)}
    return d


# ------------------------------------------------------------------ fixtures: current data
_CUR = {}


def current():
    """v19 data migrated (with the §6.2 cleanups), its oracle, and the recalculated deliverable."""
    if not _CUR:
        data, wv = mg.read_v19(CFG["v19"])
        log = mg.apply_cleanups(data)
        oracles = run_oracle(data, 1000)
        c = Calc(CFG["deliverable"], 1000)
        _CUR.update(data=data, wv=wv, log=log, oracles=oracles, calc=c)
    return _CUR


# ------------------------------------------------------------------ T1..T4 (current data)
def test_T1_oracle_parity_current():
    cur = current()
    c, data, oracles = cur["calc"], cur["data"], cur["oracles"]
    assert not c.errors, c.errors
    rows_cmp, bad = 0, []
    for m in bt.MACHINES:
        n, b = compare_machine(c, m, data, 1000, oracles[m])
        rows_cmp += n
        bad += b
    bad += compare_bucket_status(c, data, 1000, oracles)
    dash, exp = read_dashboard(c), dashboard_expected(data, 1000, oracles)
    for m in exp:
        for k, v in exp[m].items():
            if dash[m][k] != v:
                bad.append(("Dashboard", m, k, dash[m][k], v))
    # LOT#s on the dashboard = column Q of the last build (or "—")
    for m in bt.MACHINES:
        last = int(c.v("Machine Dashboard", "H%d" % (5 if m == "Old P110" else 15)))
        q = c.v(m, "Q%d" % last)
        if dash[m]["LOT#s used"] != (q or "—"):
            bad.append(("Dashboard", m, "LOT#s", dash[m]["LOT#s used"], q))
    note("T1", rows_compared=rows_cmp, cells_checked="I,J,K,L,M,N,O,Q,R + helpers + 72-key ledger per row; "
                                                      "Bucket Status D/E/F/H/I; Dashboard", mismatches=len(bad))
    assert not bad, bad[:20]


def test_T1b_cached_values_are_real_recalc():
    """The cached values stored in the deliverable equal a fresh LibreOffice recalc of the same file."""
    c = current()["calc"]
    wb = openpyxl.load_workbook(CFG["deliverable"], data_only=True)
    wf = openpyxl.load_workbook(CFG["deliverable"])
    n, bad = 0, []
    for ws in wf.worksheets:
        wv = wb[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    n += 1
                    cached = wv[cell.coordinate].value
                    fresh = c.v(ws.title, cell.coordinate)
                    cached = "" if cached is None else cached
                    if isinstance(cached, datetime.datetime):
                        cached = (cached - datetime.datetime(1899, 12, 30)).total_seconds() / 86400.0
                    if isinstance(fresh, float) and isinstance(cached, (int, float)):
                        if abs(fresh - cached) > 1e-12:
                            bad.append((ws.title, cell.coordinate, cached, fresh))
                    elif cached != fresh:
                        bad.append((ws.title, cell.coordinate, cached, fresh))
    note("T1", cached_formula_cells_verified=n, cached_mismatches=len(bad))
    assert n > 400000 and not bad, bad[:10]


def test_T2_v19_totals_parity():
    cur = current()
    c, wv = cur["calc"], cur["wv"]
    bad, n = [], 0
    for m in bt.MACHINES:
        ws = wv[m]
        for r in range(5, 105):
            n += 1
            i19 = ws["I%d" % r].value
            i19 = i19 if isinstance(i19, (int, float)) else None
            j19 = ws["J%d" % r].value
            pm19 = None
            if isinstance(j19, (int, float)):
                pm19 = j19 + (ws["AB%d" % r].value or 0) + (ws["AF%d" % r].value or 0)
            if not close(num(c.v(m, "I%d" % r)) if c.v(m, "I%d" % r) != "" else None, i19, 1e-9):
                bad.append((m, r, "I", c.v(m, "I%d" % r), i19))
            if not close(num(c.v(m, "J%d" % r)) if c.v(m, "J%d" % r) != "" else None, pm19, 1e-9):
                bad.append((m, r, "powder on machine", c.v(m, "J%d" % r), pm19))
    bs = wv["Bucket Status"]
    diffs = []
    for k, key in enumerate(bt.KEYS):
        d19 = bs["D%d" % (4 + k)].value or 0
        d20 = c.v("Bucket Status", "D%d" % (4 + k))
        if abs(d19 - d20) > 1e-9:
            diffs.append((key, d19, d20))
    note("T2", rows_compared=n, row_mismatches=len(bad), bucket_status_differences=diffs)
    assert not bad, bad
    assert [d[0] for d in diffs] == ["M2A-"] and abs(diffs[0][2] - 5.05) < 1e-9, diffs


FUNNEL_TABLE = [  # §2.4: (sheet, row, build, bucket removed, leftover, [(used by build, kg)])
    ("Old P110", 7, "B-2761", "M1B-", 3.580, [("B-2761", 3.580)]),
    ("Old P110", 11, "B-2766", "M1D-", 0.630, [("B-2766", 0.630)]),
    ("Old P110", 18, "B-2785", "M3A-", 1.770, [("B-2785", 1.770)]),
    ("New P110", 7, "B-2759", "M1F-", 2.790, [("B-2759", 2.790)]),
    ("New P110", 8, "B-2760", "M1G-", 3.670, [("B-2760", 2.760), ("B-2764", 0.910)]),
    ("New P110", 10, "B-2765", "M1H-", 2.720, [("B-2765", 2.720)]),
    ("New P110", 17, "B-2777", "M2H-", 2.400, [("B-2777", 2.400)]),
    ("New P110", 19, "B-2783", "M3F-", 1.320, [("B-2783", 1.320)]),
    ("New P110", 21, "B-2788", "M3H-", 1.670, [("B-2788", 1.670)]),
    ("New P110", 22, "B-2793", "M3D-", 3.290, [("B-2793", 3.210), ("B-2795", 0.080)]),
]

QA_TABLE = [  # §6.3
    ("Old P110", "B-2775", "M2C-20260917", "M2C-20260917, M2B-20260917"),
    ("Old P110", "B-2778", "M2B-20260917, M2C-20260917, M2D-20260917", "M2C-20260917, M2D-20260917"),
    ("Old P110", "B-2781", "M3A-20261002", "M2D-20260917"),
    ("Old P110", "B-2785", "M3B-20261002", "M3A-20261002, M3B-20261002"),
    ("Old P110", "B-2789", "M2D-20260917, M3B-20261002, M3C-20261002", "M3B-20261002, M3C-20261002"),
    ("Old P110", "B-2792", "M4A-20261009", "M4A-20261009, M3C-20261002"),
    ("Old P110", "B-2792(2nd)", "M3C-20261002, M4A-20261009, M4B-20261009", "M4A-20261009, M4B-20261009"),
    ("New P110", "B-2770", "M2G-20260917", "M2F-20260917"),
    ("New P110", "B-2773", "M2F-20260917", "M2G-20260917"),
    ("New P110", "B-2776", "M2F-20260917, M2H-20260917", "M2G-20260917, M2H-20260917"),
]


def funnel_events(c, sheet, nrows):
    """Funnel lots read back from the workbook, followed by identity through the queue.
    Returns ([{row, build, key, kg, used: [(build, kg)]}], identities of the lots still in the funnel at the end)."""
    last = bt.FIRST + nrows - 1
    ids, lots = [], {}
    for r in range(bt.FIRST, last + 1):
        cur = ids[:int(c.m(sheet, "n0", r))]
        for x in ("A", "B"):
            if c.m(sheet, "a" + x, r) == 1:
                lid = (r, c.m(sheet, "R" + x, r))
                lots[lid] = dict(row=r, build=c.v(sheet, "A%d" % r), key=lid[1], kg=c.m(sheet, "p" + x, r), used=[])
                cur.append(lid)
        for j, lid in enumerate(cur[:4]):
            t = c.m(sheet, "tk%d" % (j + 1), r)
            if t > EPS:
                lots[lid]["used"].append((c.v(sheet, "A%d" % r), t))
        ids = cur[int(c.m(sheet, "drop", r)):4]
    return sorted(lots.values(), key=lambda d: d["row"]), ids


def test_T3_expected_values():
    cur = current()
    c, data = cur["calc"], cur["data"]
    dash = read_dashboard(c)
    exp = {
        "Old P110": {"Last build #": "B-2792(3rd)", "Print date": 46309.0, "Powder used (this build)": "— (usage not entered)",
                     "LOT#s used": "—", "Slot 1 bucket + remaining": "4.980 kg M4C-",
                     "Slot 2 bucket + remaining": "3.680 kg M4B-", "Funnel contents": "0.000 kg",
                     "Powder on machine": "8.660 kg"},
        "New P110": {"Last build #": "B-2796", "Print date": 46309.0, "Powder used (this build)": "3.750 kg",
                     "LOT#s used": "M4E-20261009, M4F-20261009", "Slot 1 bucket + remaining": "0.000 kg M4E-",
                     "Slot 2 bucket + remaining": "4.150 kg M4F-", "Funnel contents": "0.000 kg",
                     "Powder on machine": "4.150 kg"},
    }
    assert datetime.date(1899, 12, 30) + datetime.timedelta(days=46309) == datetime.date(2026, 10, 14)
    for m in exp:
        for k, v in exp[m].items():
            assert dash[m][k] == v, (m, k, dash[m][k], v)
    bs = {c.v("Bucket Status", "A%d" % r): [c.v("Bucket Status", "%s%d" % (x, r)) for x in "BDEF"] for r in range(4, 76)}
    assert abs(bs["M2A-"][1] - 5.05) < 1e-9 and bs["M2A-"][2] == 0 and bs["M2A-"][3] == "Empty", bs["M2A-"]
    open_ = {"M4B-": 3.680, "M4C-": 4.980, "M4D-": 5.030, "M4F-": 4.150, "M4G-": 5.070, "M4H-": 5.065}
    for k, v in bs.items():
        if k in open_:
            assert abs(v[2] - open_[k]) < 1e-9 and v[3] == "OK", (k, v)
        elif k[:2] in ("M1", "M2", "M3", "M4") and v[0] != "":
            assert v[3] == "Empty", (k, v)
    # §2.4 funnel-event table
    got = []
    for m in bt.MACHINES:
        ev, _ = funnel_events(c, m, 1000)
        for e in ev:
            got.append((m, e["row"], e["build"], e["key"], round(e["kg"], 3), [(b, round(a, 3)) for b, a in e["used"]]))
    exp_t = [(a, b, c_, d, e, f) for a, b, c_, d, e, f in FUNNEL_TABLE]
    assert got == exp_t, (got, exp_t)
    # §6.3 table regenerated from the oracle (draw-record order) + v19 cached values
    qa = mg.qa_table(data, cur["wv"], cur["oracles"])
    got_qa = [(m, b, a, v) for m, b, r, a, v in qa]
    assert got_qa == QA_TABLE, got_qa
    # ...and column Q holds the same LOT#s (consumption order) for those builds
    for m, b, r, a, v in qa:
        q = c.v(m, "Q%d" % r)
        assert set(q.split(", ")) == set(v.split(", ")), (m, b, q, v)
    # ...and the in-workbook Change Log carries the table
    wb = openpyxl.load_workbook(CFG["deliverable"], read_only=True)
    cl = [tuple(x) for x in wb["Change Log"].iter_rows(values_only=True)]
    for m, b, a, v in QA_TABLE:
        assert any(row[1] == b and row[2] == a and row[3] == v and str(row[0]).startswith(m) for row in cl), (m, b)
    assert any(isinstance(x, str) and "QA REVIEW REQUIRED" in x for row in cl for x in row)
    note("T3", dashboard_fields=16, funnel_events=len(got), qa_rows=len(got_qa))


def test_T4_checks_current_data():
    c = current()["calc"]
    items = bt.checks_list(1000)
    rows = [(c.v("Checks", "A%d" % (bt.CHECK_ROW0 + i)), c.v("Checks", "B%d" % (bt.CHECK_ROW0 + i)),
             c.v("Checks", "C%d" % (bt.CHECK_ROW0 + i)), c.v("Checks", "D%d" % (bt.CHECK_ROW0 + i)))
            for i in range(len(items))]
    errs = sum(r[2] for r in rows if r[0] == "ERROR")
    nonzero = {r[1]: (r[2], r[3]) for r in rows if r[2]}
    assert errs == 0, nonzero
    assert nonzero == {
        "Build date missing / not a date / out of order": (2.0, "Old P110 row 5: Build # has no Date"),
        "Duplicate Build # on a sheet": (2.0, "Old P110 row 9: Duplicate Build # on this sheet"),
        "Pending build (Build # but no usage)": (1.0, "Old P110 row 22: Pending: usage not entered"),
    }, nonzero
    # the second missing-date row is New P110 row 5; duplicates are Old rows 9 (B-2762) and 11 (B-2766)
    assert c.v("New P110", "R5") == "WARN: Build # has no Date"
    assert c.v("Old P110", "A9") == "B-2762" and c.v("Old P110", "A11") == "B-2766"
    assert c.v("Old P110", "R11") == "WARN: Duplicate Build # on this sheet"
    assert c.v("Old P110", "R22") == "PENDING"
    assert c.v("Checks", "A3") == "0 ERRORS / 5 WARNINGS" and c.v("Machine Dashboard", "A3") == "0 ERRORS / 5 WARNINGS"
    note("T4", banner=c.v("Checks", "A3"), warnings={k: v[0] for k, v in nonzero.items()})


# ------------------------------------------------------------------ fuzz
OLD_CYC, NEW_CYC = [1, 2, 3, 4], [5, 6, 7, 8, 9]


def fuzz_data(rng, n_old, n_new):
    buckets = {k: {"B": round(rng.uniform(4.5, 5.3), 3), "C": datetime.datetime(2020, 1, 1), "G": None} for k in bt.KEYS}
    return {"machines": {"Old P110": gen_seq(rng, n_old, OLD_CYC), "New P110": gen_seq(rng, n_new, NEW_CYC)},
            "buckets": buckets, "settings": dict(bt.SETTINGS_DEFAULTS)}


def _pad(rng, v):
    if v is None or rng.random() > 0.05:
        return v
    return rng.choice([" %s " % v.lower(), v.lower(), "%s " % v, " " + v])


def gen_seq(rng, n, cycles):
    rows = []
    cyc, bkt = {"1": "", "2": ""}, {"1": "", "2": ""}
    d0 = datetime.datetime(2026, 1, 1)
    for i in range(n):
        if rng.random() < 0.08 and i > 0:
            rows.append({c: None for c in "ABCDEFGHP"})
            continue
        row = {"A": "F-%d" % i, "B": d0 + datetime.timedelta(days=i), "P": None,
               "C": rng.choices([None, "N", "Y"], [70, 15, 15])[0],
               "H": round(rng.uniform(0, 7), 2) if rng.random() > 0.06 else None}
        new = {}
        for s in "12":
            a = rng.choices(["none", "load", "letter", "cycle", "empty"], [55, 25, 8, 4, 8])[0]
            c = b = None
            if a == "load":
                c, b = "M%d" % rng.choice(cycles), rng.choice("ABCDEFGH")
            elif a == "letter" and cyc[s]:
                b = rng.choice("ABCDEFGH")
            elif a == "cycle" and bkt[s]:
                c = "M%d" % rng.choice(cycles)
            elif a == "empty":
                c = "EMPTY"
            new[s] = (c, b)
        # effective keys after this row; never the same bucket in both slots
        def eff(s, c, b):
            if c == "EMPTY":
                return "", ""
            return (c or cyc[s]), (b or bkt[s])
        e1, e2 = eff("1", *new["1"]), eff("2", *new["2"])
        k1 = e1[0] + e1[1] if e1[0] and e1[1] else ""
        k2 = e2[0] + e2[1] if e2[0] and e2[1] else ""
        if k1 and k1 == k2:
            new["2"] = ("EMPTY", None)
            e2 = ("", "")
        cyc["1"], bkt["1"] = e1
        cyc["2"], bkt["2"] = e2
        row["D"], row["E"] = _pad(rng, new["1"][0]), _pad(rng, new["1"][1])
        row["F"], row["G"] = _pad(rng, new["2"][0]), _pad(rng, new["2"][1])
        rows.append(row)
    return rows


def mirror(data):
    d = json.loads(json.dumps(data, default=str))
    d["buckets"] = data["buckets"]
    for m in bt.MACHINES:
        d["machines"][m] = [dict(r, D=r["F"], E=r["G"], F=r["D"], G=r["E"]) for r in data["machines"][m]]
        for r, src in zip(d["machines"][m], data["machines"][m]):
            r["B"] = src["B"]
    return d


FUZZ_NROWS = 80
_FUZZ = []


def fuzz_cases():
    if not _FUZZ:
        rng = random.Random(CFG["seed"])
        for _ in range((CFG["fuzz"] + 1) // 2):
            _FUZZ.append(fuzz_data(rng, rng.randint(5, FUZZ_NROWS), rng.randint(5, FUZZ_NROWS)))
    return _FUZZ


def test_T5_differential_fuzz():
    seqs = rows = truncated = 0
    bad = []
    stats = {"overdraw": 0, "funnel_full": 0, "in_funnel": 0, "lot5": 0, "swaps_both": 0}
    t0 = time.time()
    for k, data in enumerate(fuzz_cases()):
        oracles = run_oracle(data, FUZZ_NROWS)
        c = build_calc(data, FUZZ_NROWS, "fuzz")
        if c.errors:
            bad.append(("formula errors", k, c.errors[:3]))
        for m in bt.MACHINES:
            n, b = compare_machine(c, m, data, FUZZ_NROWS, oracles[m])
            seqs += 1
            rows += len(data["machines"][m])
            truncated += n < FUZZ_NROWS
            bad += [("case %d" % k,) + x for x in b]
            prev = {"k1": "", "k2": ""}
            for o in oracles[m][0]:     # rows where both slots changed bucket on the same row
                stats["swaps_both"] += o["k1"] != prev["k1"] and o["k2"] != prev["k2"]
                prev = o
            for _, kind, t in oracles[m][2]:
                stats["overdraw"] += t.startswith("Overdraw")
                stats["funnel_full"] += t.startswith("Funnel full")
                stats["in_funnel"] += "still in the funnel" in t
                stats["lot5"] += t.startswith("More than")
        if not any(x[0].startswith("More than") for m in bt.MACHINES for x in [(t,) for _, _, t in oracles[m][2]]):
            bad += [("case %d" % k,) + x for x in compare_bucket_status(c, data, FUZZ_NROWS, oracles)]
        if len(bad) > 50:
            break
    note("T5", seed=CFG["seed"], sequences=seqs, input_rows=rows, sequences_truncated_at_5th_funnel_lot=truncated,
         oracle_issue_counts=stats, mismatches=len(bad), seconds=round(time.time() - t0))
    assert seqs >= CFG["fuzz"] and not bad, bad[:20]


def _sym_signature(c, data, nrows):
    sig = {}
    for m in bt.MACHINES:
        for r in range(bt.FIRST, bt.FIRST + nrows):
            for col in ("I", "J", "R"):
                sig[(m, r, col)] = c.v(m, "%s%d" % (col, r))
            q = c.v(m, "Q%d" % r)
            sig[(m, r, "Qset")] = tuple(sorted(q.split(", "))) if q else ()
            for k in range(72):
                sig[(m, r, "LED%d" % k)] = c.led(m, k, r)
            sig[(m, r, "msgs")] = tuple(sorted({c.m(m, n, r) for n, _, _, _ in bt.ERR + bt.WARN} - {""}))
    for r in range(4, 76):
        for col in "DEFHI":
            v = c.v("Bucket Status", "%s%d" % (col, r))
            sig[("BS", r, col)] = v.replace("slot 1", "slot X").replace("slot 2", "slot X") if col == "I" else v
    return sig


def _sym_diff(a, b):
    out = []
    for k in a:
        x, y = a[k], b[k]
        if isinstance(x, float) and isinstance(y, float):
            if abs(x - y) > 1e-9:
                out.append((k, x, y))
        elif x != y:
            out.append((k, x, y))
    return out


def test_T6_slot_symmetry():
    n_seq, bad = 0, []
    cases = fuzz_cases()[: (CFG["sym"] + 1) // 2]
    for k, data in enumerate(cases):
        a = build_calc(data, FUZZ_NROWS, "symA")
        b = build_calc(mirror(data), FUZZ_NROWS, "symB")
        d = _sym_diff(_sym_signature(a, data, FUZZ_NROWS), _sym_signature(b, data, FUZZ_NROWS))
        n_seq += 2
        bad += [("case %d" % k,) + x for x in d]
        if len(bad) > 50:
            break
    # targeted cases
    targeted = []
    base_b = {k: {"B": 5.0, "C": datetime.datetime(2026, 1, 1), "G": None} for k in bt.KEYS}
    d1 = datetime.datetime(2026, 2, 1)
    def mk(rows):
        return {"machines": {"Old P110": rows, "New P110": []}, "buckets": base_b, "settings": dict(bt.SETTINGS_DEFAULTS)}
    def R(a, d=None, e=None, f=None, g=None, h=None, day=0):
        return {"A": a, "B": d1 + datetime.timedelta(days=day), "C": None, "D": d, "E": e, "F": f, "G": g, "H": h, "P": None}
    scen = {
        "one bucket slot1 vs slot2": ([R("b1", "M1", "A", h=1.5), R("b2", h=1.0, day=1)],
                                      [R("b1", f="M1", g="A", h=1.5), R("b2", h=1.0, day=1)]),
        "new bucket beside half-used": ([R("b1", "M1", "A", h=2.5), R("b2", f="M1", g="B", h=3.0, day=1)],
                                        [R("b1", f="M1", g="A", h=2.5), R("b2", "M1", "B", h=3.0, day=1)]),
        "both loaded same row": ([R("b1", "M1", "C", "M1", "B", h=6.0)], [R("b1", "M1", "B", "M1", "C", h=6.0)]),
    }
    for name, (ra, rb) in scen.items():
        A, B = build_calc(mk(ra), 10, "t6"), build_calc(mk(rb), 10, "t6")
        led = lambda c, key: [c.led("Old P110", bt.KEYS.index(key), r) for r in range(5, 7)]  # noqa: E731
        for key in ("M1A-", "M1B-", "M1C-"):
            if any(abs(x - y) > 1e-9 for x, y in zip(led(A, key), led(B, key))):
                targeted.append((name, key, led(A, key), led(B, key)))
        if name == "new bucket beside half-used":
            # half-used M1A- (2.5 kg left) drains first in either slot: 2.5 from M1A-, 0.5 from M1B-
            for c in (A, B):
                assert abs(c.led("Old P110", bt.KEYS.index("M1A-"), 6)) < 1e-9
                assert abs(c.led("Old P110", bt.KEYS.index("M1B-"), 6) - 4.5) < 1e-9
        if name == "both loaded same row":
            for c in (A, B):   # key order decides: M1B- first (5.0), then M1C- (1.0)
                assert abs(c.led("Old P110", bt.KEYS.index("M1B-"), 5)) < 1e-9
                assert abs(c.led("Old P110", bt.KEYS.index("M1C-"), 5) - 4.0) < 1e-9
    note("T6", fuzz_sequences_mirrored=n_seq, targeted_cases=len(scen), differences=len(bad) + len(targeted))
    assert n_seq >= CFG["sym"] and not bad and not targeted, (bad[:10], targeted)


# ------------------------------------------------------------------ T7 bad input
def _base():
    b = {k: {"B": 5.0, "C": datetime.datetime(2026, 1, 1), "G": None} for k in bt.KEYS}
    b["M9H-"] = {"B": None, "C": None, "G": None}       # Not filled
    b["M9G-"] = {"B": 5.0, "C": datetime.datetime(2026, 3, 1), "G": None}
    d = datetime.datetime(2026, 2, 1)
    def R(a, day, dd=None, e=None, f=None, g=None, h=None, cl=None):
        return {"A": a, "B": d + datetime.timedelta(days=day), "C": cl, "D": dd, "E": e, "F": f, "G": g, "H": h, "P": None}
    old = [R("b1", 0, "M1", "A", "M1", "B", 2.0), R("b2", 1, h=3.0), R("b3", 2, "M1", "C", h=2.5), R("b4", 3, h=1.0),
           R("b5", 4, h=1.0), R("b6", 5, h=1.5)]
    new = [R("n1", 0, "M5", "A", h=1.0), R("n2", 1, h=1.0)]
    return {"machines": {"Old P110": old, "New P110": new}, "buckets": b, "settings": dict(bt.SETTINGS_DEFAULTS)}, R


def _set(data, sheet, i, **kw):
    import copy
    d = copy.deepcopy(data)
    while len(d["machines"][sheet]) <= i:
        d["machines"][sheet].append({c: None for c in "ABCDEFGHP"})
    d["machines"][sheet][i].update(kw)
    return d


def test_T7_bad_input():
    base, R = _base()
    NR = 12
    ref = build_calc(base, NR, "t7base")
    assert ref.v("Checks", "A3") == "ALL CHECKS PASS", ref.v("Checks", "A3")
    blank = {c: None for c in "ABCDEFGHP"}
    def unprotect_and(fn):
        def mut(wb):
            for ws in wb.worksheets:
                ws.protection.sheet = False
            fn(wb)
        return mut
    cases = [
        # (name, data, mutate, sheet, row, expected R prefix/substring, check label, isolated)
        ("whitespace", _set(base, "Old P110", 3, F="   "), None, "Old P110", 8, "ERROR: Whitespace-only", "Whitespace-only / non-printing input", True),
        ("NBSP-only", _set(base, "Old P110", 3, D="  "), None, "Old P110", 8, "ERROR: Whitespace-only", "Whitespace-only / non-printing input", True),
        ("whitespace Build #", _set(base, "Old P110", 3, A="  "), None, "Old P110", 8, "ERROR: Whitespace-only", "Whitespace-only / non-printing input", False),
        ("M10", _set(base, "Old P110", 2, D="M10"), None, "Old P110", 7, "ERROR: Invalid entry", "Invalid Cycle / Bucket entry", False),
        ("letter I", _set(base, "Old P110", 2, E="I"), None, "Old P110", 7, "ERROR: Invalid entry", "Invalid Cycle / Bucket entry", False),
        ("EMPTY+letter", _set(base, "Old P110", 3, F="EMPTY", G="B"), None, "Old P110", 8, "ERROR: EMPTY entered", "EMPTY entered with a bucket letter", False),
        ("cycle never completed", _set(base, "New P110", 1, F="M6"), None, "New P110", 6, "ERROR: Key can't form", "Bucket key can't form", True),
        ("text usage", _set(base, "Old P110", 3, H="3.5"), None, "Old P110", 8, "ERROR: Usage must be a number", "Usage non-numeric or negative", False),
        ("negative usage", _set(base, "Old P110", 3, H=-1.0), None, "Old P110", 8, "ERROR: Usage must be a number", "Usage non-numeric or negative", False),
        ("Cleaned = X", _set(base, "Old P110", 3, C="X"), None, "Old P110", 8, "ERROR: Cleaned Chamber", "Cleaned Chamber not Y/N/blank", True),
        ("usage, nothing loaded", _set(_set(base, "New P110", 0, D=None, E=None), "New P110", 1, H=1.0), None, "New P110", 5,
         "ERROR: 1.000 kg used with no bucket loaded", "Usage with no bucket and no funnel powder", False),
        ("entries on row with no Build #", _set(base, "Old P110", 6, **dict(blank, H=1.0)), None, "Old P110", 11, "ERROR: Entries on a row with no Build #", "Slot/usage entries on a row with no Build #", True),
        ("bucket on both machines", _set(base, "New P110", 1, F="M1", G="C"), None, "New P110", 6, "ERROR: M1C- is used on both machines", "Bucket used on both machines", False),
        ("same bucket both slots", _set(base, "Old P110", 3, F="M1", G="C"), None, "Old P110", 8, "ERROR: Same bucket M1C- in both slots", "Same bucket in both slots", False),
        ("overdraw", _set(base, "Old P110", 5, H=12.0), None, "Old P110", 10, "ERROR: Overdraw: M1C- goes negative", "Overdraw (bucket goes negative)", False),
        ("typed constant over formula", base, unprotect_and(lambda wb: wb["Old P110"].__setitem__("J7", 4.78)), "Old P110", 7,
         "ERROR: Formula overwritten", "Formula overwritten on a machine-sheet row", True),
        ("typed constant over Bucket Status D", base, unprotect_and(lambda wb: wb["Bucket Status"].__setitem__("D12", 5.09)), None, None,
         None, "Bucket Status formula overwritten (D/E/F/H/I or hidden helpers)", True),
        ("key not filled", _set(base, "New P110", 1, F="M9", G="H"), None, "New P110", 6, "ERROR: M9H- is not in Bucket Status or is Not filled", "Bucket not in Bucket Status / Not filled", False),
        ("loaded before fill date", _set(base, "New P110", 1, F="M9", G="G"), None, "New P110", 6, "ERROR: M9G- loaded before its Date filled", "Bucket loaded before its Date filled (or no Date filled)", False),
        ("reload bucket in funnel", _set(_set(base, "Old P110", 3, F="M1", G="D", H=0.5), "Old P110", 4, F="M1", G="B"), None, "Old P110", 9,
         "ERROR: M1B- loaded while its leftover is still in the funnel", "Bucket loaded while its leftover is in the funnel", False),
        ("funnel overflow", _set(_set(base, "Old P110", 3, F="M2", G="A", H=0.0), "Old P110", 4, D="M2", E="B", H=0.0), None, "Old P110", 9,
         "WARN: Funnel full: ", "Funnel full (overflow kept on bucket)", False),
    ]
    # 5 funnel lots: four removals with no builds consuming, then a fifth
    five = base
    for i, (c1, b1) in enumerate([("M2", "A"), ("M2", "B"), ("M2", "C"), ("M2", "D"), ("M2", "E")]):
        five = _set(five, "Old P110", 6 + i, A="x%d" % i, B=datetime.datetime(2026, 2, 10 + i), D=c1, E=b1, H=0.0)
    s = dict(bt.SETTINGS_DEFAULTS)
    s["B6"] = 50
    five["settings"] = s
    cases.append(("5 funnel lots", five, None, "Old P110", 15, "ERROR: More than 4 bucket leftovers in funnel", "More than 4 funnel lots", False))
    results, fails = [], []
    labels = {lab: i for i, (_, lab, _, _, _) in enumerate(bt.checks_list(NR))}
    def first_changed(data):
        for m in bt.MACHINES:
            a, b = data["machines"][m], base["machines"][m]
            for i in range(max(len(a), len(b))):
                if (a[i] if i < len(a) else None) != (b[i] if i < len(b) else None):
                    return m, bt.FIRST + i
        return None, 10 ** 6
    for name, data, mutate, sheet, row, rexp, label, isolated in cases:
        c = build_calc(data, NR, "t7", mutate)
        msheet, mrow = first_changed(data)
        ok = not c.errors
        if sheet:
            got = c.v(sheet, "R%d" % row)
            ok &= got.startswith(rexp)
        cnt = c.v("Checks", "C%d" % (bt.CHECK_ROW0 + labels[label]))
        ok &= isinstance(cnt, float) and cnt >= 1
        banner = c.v("Checks", "A3")
        ok &= banner != "ALL CHECKS PASS"
        # other rows unaffected: every row before the bad row; all other rows when the input can't change state
        for m in bt.MACHINES:
            for r in range(bt.FIRST, bt.FIRST + NR):
                if sheet is None and mutate is None:
                    continue
                if m == msheet and (r == mrow or (r > mrow and not isolated)):
                    continue
                if mutate is not None and m == sheet and r == row:
                    continue
                for col in ("I", "J"):
                    if c.v(m, "%s%d" % (col, r)) != ref.v(m, "%s%d" % (col, r)) and not (name.startswith("typed") and col == "J" and r == 7):
                        ok = False
                        fails.append((name, m, r, col, c.v(m, "%s%d" % (col, r)), ref.v(m, "%s%d" % (col, r))))
        results.append((name, ok, c.v(sheet, "R%d" % row) if sheet else "", cnt, banner))
        if not ok and not any(f[0] == name for f in fails):
            fails.append((name, c.v(sheet, "R%d" % row) if sheet else None, cnt, banner, c.errors))
    # must pass: padded / lower-case entries normalize to the same result as clean ones
    padded = _set(_set(base, "Old P110", 2, D=" m1 ", E=" c "), "Old P110", 0, D="m1", E=" a", F=" M1", G="b ")
    p = build_calc(padded, NR, "t7pad")
    same = all(p.v(m, "%s%d" % (col, r)) == ref.v(m, "%s%d" % (col, r))
               for m in bt.MACHINES for r in range(5, 5 + NR) for col in "IJKLMNOQR")
    results.append(("' m1 ' / ' c ' / NBSP-padded pass", same and p.v("Checks", "A3") == "ALL CHECKS PASS", "", 0, p.v("Checks", "A3")))
    if not results[-1][1]:
        fails.append(("padded", p.v("Checks", "A3")))
    note("T7", cases=[(n, "pass" if ok else "FAIL", r) for n, ok, r, _, _ in results])
    assert all(ok for _, ok, _, _, _ in results) and not fails, (fails, [r for r in results if not r[1]])


# ------------------------------------------------------------------ T8 edge cases
def test_T8_edge_cases():
    b = {k: {"B": 5.0, "C": datetime.datetime(2026, 1, 1), "G": None} for k in bt.KEYS}
    d = datetime.datetime(2026, 2, 1)
    def R(a, day, dd=None, e=None, f=None, g=None, h=None, cl=None):
        return {"A": a, "B": d + datetime.timedelta(days=day), "C": cl, "D": dd, "E": e, "F": f, "G": g, "H": h, "P": None}
    blank = {c: None for c in "ABCDEFGHP"}
    def mk(old, new=None, nr=12):
        return {"machines": {"Old P110": old, "New P110": new or []}, "buckets": b, "settings": dict(bt.SETTINGS_DEFAULTS)}
    checks = {}
    # first row + blank gap rows: identical results with and without gaps (state carries through)
    rows = [R("b1", 0, "M1", "A", "M1", "B", 2.0), R("b2", 1, "M1", "C", h=1.0), R("b3", 2, h=4.0)]
    gapped = [rows[0], dict(blank), dict(blank), rows[1], dict(blank), rows[2]]
    A, G = build_calc(mk(rows), 12, "t8"), build_calc(mk(gapped), 12, "t8")
    checks["first row computes"] = A.v("Old P110", "I5") == 2.0 and A.v("Old P110", "L5") == "2.000 kg M1A-"
    checks["blank gap rows carry state"] = all(A.v("Old P110", "%s%d" % (c, ra)) == G.v("Old P110", "%s%d" % (c, rg))
                                               for ra, rg in ((5, 5), (6, 8), (7, 10)) for c in "IJKLMNOQ") \
        and G.v("Old P110", "R6") == "" and G.v("Checks", "A3") == "ALL CHECKS PASS"
    # funnel lot spanning 3+ builds: M1A- leftover 3.0 kg used over b2..b4
    span = [R("b1", 0, "M1", "A", h=2.0), R("b2", 1, "M1", "B", h=1.0), R("b3", 2, h=1.0), R("b4", 3, h=1.5)]
    S = build_calc(mk(span), 12, "t8")
    ev, _ = funnel_events(S, "Old P110", 12)
    checks["funnel lot spans 3 builds"] = len(ev) == 1 and [x[0] for x in ev[0]["used"]] == ["b2", "b3", "b4"] \
        and abs(sum(x[1] for x in ev[0]["used"]) - 3.0) < 1e-9
    # pending build, then usage added later
    pend = [R("b1", 0, "M1", "A", h=2.0), R("b2", 1, h=None)]
    P1 = build_calc(mk(pend), 12, "t8")
    pend[1]["H"] = 1.25
    P2 = build_calc(mk(pend), 12, "t8")
    checks["pending then usage added"] = P1.v("Old P110", "R6") == "PENDING" and P1.v("Old P110", "I6") == "" \
        and P1.v("Old P110", "N6") == "3.000 kg M1A-" and P2.v("Old P110", "R6") == "OK" \
        and P2.v("Old P110", "L6") == "1.250 kg M1A-" and abs(P2.v("Old P110", "J6") - 1.75) < 1e-9
    # usage exactly equal to what's available -> J = 0, no error
    ex = [R("b1", 0, "M1", "A", h=2.0), R("b2", 1, h=3.0)]
    E = build_calc(mk(ex), 12, "t8")
    checks["usage == available: J = 0, OK"] = abs(E.v("Old P110", "J6")) < 1e-12 and E.v("Old P110", "R6") == "OK" \
        and E.v("Checks", "A3") == "ALL CHECKS PASS"
    # Cleaned Chamber = Y adds 1.200 kg
    Y = build_calc(mk([R("b1", 0, "M1", "A", h=2.0, cl="Y")]), 12, "t8")
    checks["Cleaned Chamber Y adds 1.200 kg"] = abs(Y.v("Old P110", "I5") - 3.2) < 1e-12
    # M9H- (Bucket Status row 75) resolves
    M = build_calc(mk([R("b1", 0, "M9", "H", h=1.0)]), 12, "t8")
    checks["M9H- resolves"] = M.v("Old P110", "Q5") == "M9H-20260101" and abs(M.v("Bucket Status", "E75") - 4.0) < 1e-9 \
        and M.v("Bucket Status", "I75") == "Old P110 slot 1"
    # row 1004 computes (full 1000-row workbook, last row used)
    full_rows = [R("b1", 0, "M1", "A", h=1.0)] + [dict(blank) for _ in range(998)] + [R("b1000", 5, h=0.5)]
    F = build_calc(mk(full_rows), 1000, "t8full")
    checks["row 1004 computes"] = F.v("Old P110", "A1004") == "b1000" and abs(F.v("Old P110", "J1004") - 3.5) < 1e-9 \
        and F.v("Old P110", "R1004") == "OK" and F.v("Machine Dashboard", "B5") == "b1000"
    note("T8", cases={k: "pass" if v else "FAIL" for k, v in checks.items()})
    assert all(checks.values()), checks


# ------------------------------------------------------------------ T9 mass balance
def _mass_balance(c, nrows, label):
    bad = []
    last = bt.FIRST + nrows - 1
    drawn = {}
    for m in bt.MACHINES:
        si = sum(c.v(m, "I%d" % r) for r in range(bt.FIRST, last + 1) if isinstance(c.v(m, "I%d" % r), float))
        sd = sum(c.m(m, "noB", r) for r in range(bt.FIRST, last + 1))
        for r in range(bt.FIRST, last + 1):
            for j in range(1, 7):
                k, q = c.m(m, "DK%d" % j, r), c.m(m, "DQ%d" % j, r)
                if k != "":
                    sd += q
                    drawn[k] = drawn.get(k, 0.0) + q
        if abs(si - sd) > 1e-9 * max(1, nrows / 100):
            bad.append((label, m, "sum I vs draws + unsupplied", si, sd))
        ev, ids = funnel_events(c, m, nrows)
        for e in ev:
            used = sum(a for _, a in e["used"])
            left = 0.0
            for j, lid in enumerate(ids):
                if lid == (e["row"], e["key"]):
                    left = c.m(m, "FQ%d" % (j + 1), last)
            if abs(e["kg"] - used - left) > 1e-9:
                bad.append((label, m, "funnel lot", e, used, left))
    for k, key in enumerate(bt.KEYS):
        if abs(c.v("Bucket Status", "D%d" % (4 + k)) - drawn.get(key, 0.0)) > 1e-9:
            bad.append((label, "BS", key, c.v("Bucket Status", "D%d" % (4 + k)), drawn.get(key, 0.0)))
    return bad


def test_T9_mass_balance():
    bad = _mass_balance(current()["calc"], 1000, "current data")
    n = 1
    for k, data in enumerate(fuzz_cases()[:40]):
        oracles = run_oracle(data, FUZZ_NROWS)
        if any(t.startswith("More than") for m in bt.MACHINES for _, _, t in oracles[m][2]):
            continue
        bad += _mass_balance(build_calc(data, FUZZ_NROWS, "t9"), FUZZ_NROWS, "fuzz %d" % k)
        n += 1
    note("T9", workbooks=n, mismatches=len(bad))
    assert not bad, bad[:10]


# ------------------------------------------------------------------ T10 compatibility
BANNED = ["LET", "LAMBDA", "XLOOKUP", "XMATCH", "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "TEXTJOIN", "CONCAT",
          "IFS", "SWITCH", "MAXIFS", "MINIFS", "OFFSET", "INDIRECT", "TODAY", "NOW", "RAND", "RANDBETWEEN", "RANDARRAY",
          "CELL", "INFO"]


def test_T10_compatibility():
    path = CFG["deliverable"]
    z = zipfile.ZipFile(path)
    names = z.namelist()
    assert not any("vbaProject" in n for n in names), "macros present"
    fns, nform, bad = set(), 0, []
    for n in names:
        if n.startswith("xl/worksheets/sheet"):
            x = z.read(n).decode("utf8")
            if 't="array"' in x or " cm=" in x:
                bad.append((n, "array / dynamic-array formula"))
            for f in re.findall(r"<f>(.*?)</f>", x, re.S):
                nform += 1
                f2 = re.sub(r'"[^"]*"', '""', f.replace("&quot;", '"'))
                fns.update(re.findall(r"([A-Za-z_][A-Za-z0-9_.]*)\(", f2))
    used = sorted(fns)
    banned = [f for f in used if f.upper().replace("_XLFN.", "") in BANNED]
    xlfn = [f for f in used if f.upper().startswith("_XLFN.") and f.upper() not in ("_XLFN.ISFORMULA", "_XLFN.UNICHAR")]
    c = current()["calc"]
    errtxt = [(s, i) for s, arr in c.a.items() for i, row in enumerate(arr) for v in row
              if isinstance(v, str) and re.search(r"#(REF|NAME|VALUE|DIV/0|N/A|NUM|NULL)[!?]?|Err:\d{3}", v)]
    note("T10", formulas_scanned=nform, functions_used=used, banned_found=banned, error_cells=len(c.errors) + len(errtxt))
    assert not banned and not xlfn and not bad and not c.errors and not errtxt, (banned, xlfn, bad, c.errors[:5], errtxt[:5])


# ------------------------------------------------------------------ T11 performance
def test_T11_performance():
    rng = random.Random(CFG["seed"] + 11)
    data = fuzz_data(rng, 1000, 1000)
    c = build_calc(data, 1000, "t11")
    note("T11", rows_per_sheet=1000, recalc_seconds=round(c.recalc_s, 2), error_cells=len(c.errors))
    assert c.recalc_s < 10 and not c.errors, (c.recalc_s, c.errors[:5])


def test_T6b_mutant_is_detected():
    """Harness self-check: a v19-style 'slot 1 always drains first' workbook must FAIL the oracle and symmetry checks."""
    idx = [i for i, h in enumerate(bt.L.helpers) if h[0] == "f1"][0]
    orig = bt.L.helpers[idx]
    bt.L.helpers[idx] = ("f1", orig[1], '=IF(@K1<>"",1,0)')
    bt.row_formulas.cache_clear()
    try:
        cur = current()
        c = build_calc(cur["data"], 1000, "mutant")
        n_bad = sum(len(compare_machine(c, m, cur["data"], 1000, cur["oracles"][m])[1]) for m in bt.MACHINES)
        data = fuzz_cases()[0]
        d = _sym_diff(_sym_signature(build_calc(data, FUZZ_NROWS, "mA"), data, FUZZ_NROWS),
                      _sym_signature(build_calc(mirror(data), FUZZ_NROWS, "mB"), data, FUZZ_NROWS))
    finally:
        bt.L.helpers[idx] = orig
        bt.row_formulas.cache_clear()
    note("T6", mutant_slot1_first_oracle_mismatches=n_bad, mutant_slot1_first_symmetry_differences=len(d))
    assert n_bad > 0 and len(d) > 0, (n_bad, len(d))


TESTS = [test_T1_oracle_parity_current, test_T1b_cached_values_are_real_recalc, test_T2_v19_totals_parity,
         test_T3_expected_values, test_T4_checks_current_data, test_T5_differential_fuzz, test_T6_slot_symmetry,
         test_T6b_mutant_is_detected, test_T7_bad_input, test_T8_edge_cases, test_T9_mass_balance, test_T10_compatibility, test_T11_performance]

DESCR = {
    "T1": "Oracle parity on current data (every row, helpers, ledger, Bucket Status, Dashboard) + cached values = real recalc",
    "T2": "v19 totals parity (I and powder on machine on every row; Bucket Status D except M2A-)",
    "T3": "Expected values (Dashboard, Bucket Status, §2.4 funnel-event table, §6.3 QA table, Change Log)",
    "T4": "Checks on current data (0 errors; exactly the expected warnings)",
    "T5": "Differential fuzz vs oracle",
    "T6": "Slot symmetry (mirrored fuzz + targeted cases)",
    "T7": "Bad input raises the right ERROR/WARN without disturbing other rows",
    "T8": "Edge cases",
    "T9": "Mass balance (sum I = sum of draws + unsupplied usage; Bucket Status D = draws; funnel lots in = out + left)",
    "T10": "Compatibility (banned/volatile functions, macros, arrays, error cells)",
    "T11": "Performance (full recalc, 1,000 rows on both sheets)",
}


def write_report(results, total_s):
    lines = ["# P110 Powder Tracker v20 — TEST_REPORT", "",
             "Generated by `tests/test_tracker.py` on %s. LibreOffice headless recalculation (UNO `calculateAll`), values read back "
             "and compared with `ref_model.py`. Python %s, openpyxl %s." % (
                 datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), sys.version.split()[0], openpyxl.__version__), "",
             "Workbook under test: `P110_Powder_Tracker_v20.xlsx`. Fuzz seed: **%d**. Total run time: %d s." % (CFG["seed"], total_s), "",
             "This report records automated tests. It is not a software validation record; formal validation of this "
             "spreadsheet (ISO 13485 §4.1.6) is for QA to plan and approve.", "",
             "| Test | Description | Result |", "|---|---|---|"]
    by = {}
    for name, ok, err in results:
        t = name.split("_")[1]
        by.setdefault(t, []).append((ok, err))
    for t in DESCR:
        ok = all(o for o, _ in by.get(t, [(False, "not run")]))
        lines.append("| %s | %s | %s |" % (t, DESCR[t], "PASS" if ok else "**FAIL**"))
    lines += ["", "## Details", ""]
    for t in DESCR:
        lines.append("### %s — %s" % (t, DESCR[t]))
        for k, v in REPORT.get(t, {}).items():
            if isinstance(v, list) and v and isinstance(v[0], tuple):
                lines.append("- %s:" % k)
                for item in v:
                    lines.append("  - " + " | ".join(str(x) for x in item))
            elif isinstance(v, dict):
                lines.append("- %s:" % k)
                for a, b in v.items():
                    lines.append("  - %s: %s" % (a, b))
            else:
                lines.append("- %s: %s" % (k, v))
        for ok, err in by.get(t, []):
            if not ok:
                lines.append("- **FAILURE:** `%s`" % str(err)[:1500])
        lines.append("")
    open(os.path.join(ROOT, "TEST_REPORT.md"), "w").write("\n".join(lines))


def main(argv):
    args = list(argv)
    if args and not args[0].startswith("--"):
        CFG["v19"] = args.pop(0)
    while args:
        k = args.pop(0)
        CFG[k.lstrip("-")] = int(args.pop(0)) if k in ("--seed", "--fuzz", "--sym") else args.pop(0)
    results = []
    t0 = time.time()
    only = os.environ.get("P110_ONLY")
    try:
        for fn in TESTS:
            if only and not any(o in fn.__name__ for o in only.split(",")):
                continue
            t = time.time()
            try:
                fn()
                results.append((fn.__name__, True, None))
                print("PASS %-45s %6.1fs" % (fn.__name__, time.time() - t), flush=True)
            except Exception as e:  # noqa: BLE001
                results.append((fn.__name__, False, "%s: %s" % (type(e).__name__, e)))
                print("FAIL %-45s %6.1fs\n%s" % (fn.__name__, time.time() - t, traceback.format_exc()[-3000:]), flush=True)
    finally:
        if _LO is not None:
            _LO.__exit__(None, None, None)
    if not only:
        write_report(results, time.time() - t0)
    return all(ok for _, ok, _ in results)


if __name__ == "__main__":
    sys.exit(0 if main(sys.argv[1:]) else 1)
